import os
import sys
import io
import json
import re
import uuid
import time
import threading
import webbrowser
import sqlite3
import concurrent.futures
import winrm
import zipfile
import urllib.parse
from flask import Flask, request, jsonify, render_template, send_file, Response
from flask_cors import CORS
from flask_sock import Sock
import paramiko
from cryptography.fernet import Fernet
import openpyxl
from datetime import datetime, timedelta
import subprocess
import base64
import requests

# 動態設定 templates 與 static 的路徑以支援 PyInstaller 打包
if getattr(sys, 'frozen', False):
    template_folder = os.path.join(sys._MEIPASS, 'templates')
    static_folder = os.path.join(sys._MEIPASS, 'static')
else:
    template_folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates')
    static_folder = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'static')

app = Flask(__name__, template_folder=template_folder, static_folder=static_folder)
app.config['TEMPLATES_AUTO_RELOAD'] = True
CORS(app)
sock = Sock(app)

# 取得執行檔或程式所在目錄
def get_base_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

DB_FILE_PATH = os.path.join(get_base_dir(), "monitor.db")
KEY_PATH = os.path.join(get_base_dir(), "db_monitor.key")

def init_db():
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS servers (
                id TEXT PRIMARY KEY,
                name TEXT,
                category TEXT,
                host TEXT,
                port INTEGER,
                username TEXT,
                password TEXT,
                key_path TEXT,
                os_type TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS maintenance_templates (
                id TEXT PRIMARY KEY,
                name TEXT,
                customer TEXT,
                contact TEXT,
                version TEXT,
                location TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS system_settings (
                key TEXT PRIMARY KEY,
                value TEXT
            )
        ''')
        
        # 自動為舊的資料庫補上 os_type 欄位
        try:
            cursor.execute("ALTER TABLE servers ADD COLUMN os_type TEXT")
        except sqlite3.OperationalError:
            # 若欄位已存在，會拋出 error，直接忽略即可
            pass
            
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"Database initialization failed: {e}")

def get_db_setting(key, default=""):
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT value FROM system_settings WHERE key = ?", (key,))
        row = cursor.fetchone()
        conn.close()
        if row and row[0] is not None:
            return row[0]
    except Exception as e:
        print(f"get_db_setting error for key '{key}': {e}")
    return default

def set_db_setting(key, value):
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        cursor.execute("INSERT OR REPLACE INTO system_settings (key, value) VALUES (?, ?)", (key, value))
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print(f"set_db_setting error for key '{key}': {e}")
        return False

def init_probe_test_devices():
    """確保 6 台 Probe 測試設備存在於資料庫（分類 __probe_test__），若已存在則不覆蓋。"""
    TEST_DEVICES = [
        {'id': '__probe_test_1', 'name': '1.1.1.1', 'host': '1.1.1.1'},
        {'id': '__probe_test_2', 'name': '2.2.2.2', 'host': '2.2.2.2'},
        {'id': '__probe_test_3', 'name': '3.3.3.3', 'host': '3.3.3.3'},
        {'id': '__probe_test_4', 'name': '4.4.4.4', 'host': '4.4.4.4'},
        {'id': '__probe_test_5', 'name': '5.5.5.5', 'host': '5.5.5.5'},
        {'id': '__probe_test_6', 'name': '6.6.6.6', 'host': '6.6.6.6'},
    ]
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        # 取得目前已存在的測試設備 ID
        cursor.execute("SELECT id FROM servers WHERE category = '__probe_test__'")
        existing_ids = {row[0] for row in cursor.fetchall()}
        for d in TEST_DEVICES:
            if d['id'] not in existing_ids:
                enc_pwd = encrypt_password('1111')
                cursor.execute(
                    'INSERT INTO servers (id, name, category, host, port, username, password, key_path, os_type) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)',
                    (d['id'], d['name'], '__probe_test__', d['host'], 22, 'root', enc_pwd, '', 'debian')
                )
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"init_probe_test_devices failed: {e}")

# 密碼加密與解密輔助邏輯
def get_cipher():
    if not os.path.exists(KEY_PATH):
        key = Fernet.generate_key()
        try:
            with open(KEY_PATH, "wb") as f:
                f.write(key)
        except Exception:
            pass
    else:
        try:
            with open(KEY_PATH, "rb") as f:
                key = f.read()
        except Exception:
            key = Fernet.generate_key()
    return Fernet(key)

def encrypt_password(password):
    if not password:
        return ""
    try:
        cipher = get_cipher()
        return cipher.encrypt(password.encode('utf-8')).decode('utf-8')
    except Exception:
        return password

def decrypt_password(encrypted_password):
    if not encrypted_password:
        return ""
    try:
        cipher = get_cipher()
        return cipher.decrypt(encrypted_password.encode('utf-8')).decode('utf-8')
    except Exception:
        return encrypted_password

# 讀寫設定檔的輔助函數
def read_servers():
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM servers")
        rows = cursor.fetchall()
        servers = []
        for row in rows:
            s = dict(row)
            if "password" in s and s["password"]:
                s["password"] = decrypt_password(s["password"])
            if "os_type" not in s or not s["os_type"]:
                s["os_type"] = "debian"
            servers.append(s)
        conn.close()
        return servers
    except Exception as e:
        print(f"Error reading servers from DB: {e}")
        return []

def write_servers(servers):
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM servers")
        for s in servers:
            enc_pwd = encrypt_password(s.get("password", ""))
            cursor.execute('''
                INSERT INTO servers (id, name, category, host, port, username, password, key_path, os_type)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                s.get("id"),
                s.get("name"),
                s.get("category", "預設分類"),
                s.get("host"),
                int(s.get("port", 22)),
                s.get("username"),
                enc_pwd,
                s.get("key_path", ""),
                s.get("os_type", "debian")
              )
            )
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print(f"Error writing servers to DB: {e}")
        return False

# 解析 top 輸出的核心邏輯
def parse_top_output(text):
    # 若 top 輸出了多個輪次（例如跑了 2 次），我們只取最後一次的輸出結果
    blocks = text.split("top - ")
    if len(blocks) > 1:
        text = "top - " + blocks[-1]

    lines = text.split('\n')
    result = {
        "uptime": "未知",
        "load_avg": [0.0, 0.0, 0.0],
        "tasks": {"total": 0, "running": 0, "sleeping": 0, "stopped": 0, "zombie": 0},
        "cpu": {"user": 0.0, "system": 0.0, "idle": 0.0, "wait": 0.0},
        "mem": {"total": 0.0, "used": 0.0, "free": 0.0, "cache": 0.0, "unit": "MiB"},
        "swap": {"total": 0.0, "used": 0.0, "free": 0.0, "avail": 0.0, "unit": "MiB"},
        "processes": []
    }
    
    process_started = False
    headers = []
    
    for line in lines:
        line_stripped = line.strip()
        if not line_stripped:
            continue
            
        # 1. 解析 Uptime 與 Load Average
        if "load average:" in line_stripped:
            # 格式例如: top - 17:22:20 up 10 days,  2:30,  1 user,  load average: 0.00, 0.01, 0.05
            up_match = re.search(r'up\s+(.*?),\s*\d+\s+user', line_stripped)
            if not up_match:
                up_match = re.search(r'up\s+(.*?),\s*\d+\s+users', line_stripped)
            if up_match:
                result["uptime"] = up_match.group(1)
            
            load_match = re.search(r'load average:\s*([0-9.,]+),\s*([0-9.,]+),\s*([0-9.,]+)', line_stripped)
            if load_match:
                result["load_avg"] = [
                    float(load_match.group(1).replace(',', '.')),
                    float(load_match.group(2).replace(',', '.')),
                    float(load_match.group(3).replace(',', '.'))
                ]
                
        # 2. 解析 Tasks
        elif line_stripped.startswith("Tasks:"):
            # Tasks: 120 total,   1 running, 119 sleeping,   0 stopped,   0 zombie
            tasks_match = re.findall(r'(\d+)\s+([a-zA-Z\s]+)', line_stripped)
            for val, key in tasks_match:
                key_clean = key.strip().lower()
                val_i = int(val)
                if "total" in key_clean: result["tasks"]["total"] = val_i
                elif "running" in key_clean or "run" in key_clean: result["tasks"]["running"] = val_i
                elif "sleeping" in key_clean or "sleep" in key_clean: result["tasks"]["sleeping"] = val_i
                elif "stopped" in key_clean or "stop" in key_clean: result["tasks"]["stopped"] = val_i
                elif "zombie" in key_clean or "zomb" in key_clean: result["tasks"]["zombie"] = val_i
                    
        # 3. 解析 CPU
        elif "%Cpu" in line_stripped:
            # %Cpu(s):  0.3 us,  0.3 sy,  0.0 ni, 99.3 id,  0.0 wa,  0.0 hi,  0.0 si,  0.0 st
            cpu_match = re.findall(r'([0-9]+[.,]?[0-9]*)\s+([a-zA-Z]+)', line_stripped)
            for val, key in cpu_match:
                try:
                    val_f = float(val.replace(',', '.'))
                    if key == "us": result["cpu"]["user"] = val_f
                    elif key == "sy": result["cpu"]["system"] = val_f
                    elif key == "id": result["cpu"]["idle"] = val_f
                    elif key == "wa": result["cpu"]["wait"] = val_f
                except Exception:
                    pass

        # 4. 解析 記憶體 (Mem)
        elif ("Mem" in line_stripped.split(':')[0] or "mem" in line_stripped.split(':')[0]) and ":" in line_stripped:
            # MiB Mem :   1988.2 total,    120.5 free,    980.2 used,    887.5 buff/cache
            unit_match = re.match(r'^([a-zA-Z]+)\s+Mem', line_stripped, re.IGNORECASE)
            if unit_match:
                result["mem"]["unit"] = unit_match.group(1)
            
            parts = re.findall(r'([0-9]+[.,]?[0-9]*)\s+([a-zA-Z/]+)', line_stripped)
            for val, key in parts:
                try:
                    val_f = float(val.replace(',', '.'))
                    key_clean = key.lower()
                    if "total" in key_clean: result["mem"]["total"] = val_f
                    elif "free" in key_clean: result["mem"]["free"] = val_f
                    elif "used" in key_clean: result["mem"]["used"] = val_f
                    elif "buff" in key_clean or "cache" in key_clean: result["mem"]["cache"] = val_f
                except Exception:
                    pass
                
        # 5. 解析 Swap
        elif ("Swap" in line_stripped.split(':')[0] or "swap" in line_stripped.split(':')[0]) and ":" in line_stripped:
            # MiB Swap:   2048.0 total,   2048.0 free,      0.0 used.    860.1 avail Mem
            unit_match = re.match(r'^([a-zA-Z]+)\s+Swap', line_stripped, re.IGNORECASE)
            if unit_match:
                result["swap"]["unit"] = unit_match.group(1)
                
            parts = re.findall(r'([0-9]+[.,]?[0-9]*)\s+([a-zA-Z/.\s]+)', line_stripped)
            for val, key in parts:
                try:
                    val_f = float(val.replace(',', '.'))
                    key_clean = key.lower()
                    if "total" in key_clean: result["swap"]["total"] = val_f
                    elif "free" in key_clean: result["swap"]["free"] = val_f
                    elif "used" in key_clean: result["swap"]["used"] = val_f
                    elif "avail" in key_clean: result["swap"]["avail"] = val_f
                except Exception:
                    pass

        # 6. 程序列表解析起點
        elif "PID" in line_stripped and "USER" in line_stripped and "COMMAND" in line_stripped:
            process_started = True
            headers = [h.strip() for h in line_stripped.split() if h.strip()]
            continue
            
        elif process_started:
            cols = line_stripped.split()
            if len(cols) >= 8:
                try:
                    # 預設對應索引：PID USER PR NI VIRT RES SHR S %CPU %MEM TIME+ COMMAND
                    pid = cols[0]
                    user = cols[1]
                    state = cols[7]
                    cpu = float(cols[8].replace(',', '.'))
                    mem = float(cols[9].replace(',', '.'))
                    time_val = cols[10]
                    command = " ".join(cols[11:])
                    result["processes"].append({
                        "pid": pid,
                        "user": user,
                        "state": state,
                        "cpu": cpu,
                        "mem": mem,
                        "time": time_val,
                        "command": command
                    })
                except Exception:
                    pass

    return result

LAST_NET_STATS = {}

def format_speed(bytes_per_sec):
    if bytes_per_sec < 1024:
        return f"{bytes_per_sec:.1f} B/s"
    elif bytes_per_sec < 1024 * 1024:
        return f"{bytes_per_sec / 1024:.1f} KB/s"
    else:
        return f"{bytes_per_sec / (1024 * 1024):.1f} MB/s"

def parse_df_output(df_text):
    disk_info = {"total": "0G", "used": "0G", "avail": "0G", "percent": 0}
    for line in df_text.splitlines():
        line_stripped = line.strip()
        if not line_stripped:
            continue
        cols = line_stripped.split()
        if len(cols) >= 6 and cols[5] == "/":
            disk_info["total"] = cols[1]
            disk_info["used"] = cols[2]
            disk_info["avail"] = cols[3]
            try:
                disk_info["percent"] = int(cols[4].replace('%', ''))
            except Exception:
                pass
            break
    return disk_info

def parse_net_output(net_text):
    rx_bytes = 0
    tx_bytes = 0
    for line in net_text.splitlines():
        line_stripped = line.strip()
        if ":" not in line_stripped:
            continue
        parts = line_stripped.split(":")
        if len(parts) < 2:
            continue
        iface = parts[0].strip()
        # 忽略 loopback 與 docker 虛擬網卡
        if iface == "lo" or iface.startswith("docker") or iface.startswith("br-") or iface.startswith("veth"):
            continue
        cols = parts[1].split()
        if len(cols) >= 9:
            try:
                rx_bytes += int(cols[0])
                tx_bytes += int(cols[8])
            except Exception:
                pass
    return {"rx": rx_bytes, "tx": tx_bytes}

# 透過 SSH 連線並執行 top 獲取資料
def fetch_top_data(server):
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    server_id = server.get("id", "temp")
    host = server.get("host")
    port = int(server.get("port", 22))
    username = server.get("username")
    password = server.get("password")
    key_path = server.get("key_path")
    
    try:
        if key_path and os.path.exists(key_path):
            try:
                key = paramiko.RSAKey.from_private_key_file(key_path)
            except Exception:
                try:
                    key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except Exception:
                    key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(host, port=port, username=username, pkey=key, timeout=10, disabled_algorithms=None)
        else:
            ssh.connect(host, port=port, username=username, password=password, timeout=10, disabled_algorithms=None)
            
        cmd = "LC_ALL=C top -b -n 2 -d 0.2; echo \"===DF===\"; df -h; echo \"===NET===\"; cat /proc/net/dev; echo \"===SSL===\"; grep -m 1 -i '\"EnableSSL\"' /PIXIS/Probe_Daemon/Config/ProbeDaemonSetting.cfg 2>/dev/null || true"
        stdin, stdout, stderr = ssh.exec_command(cmd)
        exit_status = stdout.channel.recv_exit_status()
        if exit_status != 0:
            err_msg = stderr.read().decode('utf-8', errors='ignore')
            raise Exception(f"SSH指令執行失敗 (Exit Code {exit_status}): {err_msg}")
            
        output = stdout.read().decode('utf-8', errors='ignore')
        
        # 暫時寫入除錯檔案以檢視原始 top 輸出
        try:
            with open("debug_top.txt", "w", encoding="utf-8") as f:
                f.write(output)
        except Exception:
            pass

        # 分割不同的輸出區塊
        top_part = ""
        df_part = ""
        net_part = ""
        ssl_part = ""
        
        if "===DF===" in output:
            parts = output.split("===DF===")
            top_part = parts[0]
            rest = parts[1]
            if "===NET===" in rest:
                net_parts = rest.split("===NET===")
                df_part = net_parts[0]
                rest_net = net_parts[1]
                if "===SSL===" in rest_net:
                    ssl_split = rest_net.split("===SSL===")
                    net_part = ssl_split[0]
                    ssl_part = ssl_split[1]
                else:
                    net_part = rest_net
            else:
                df_part = rest
        else:
            top_part = output

        # 解析各個區塊
        result = parse_top_output(top_part)
        result["disk"] = parse_df_output(df_part)
        
        # 解析 Probe SSL 設定
        ssl_val = "未知"
        if ssl_part:
            m = re.search(r'"EnableSSL"\s*:\s*(true|false)', ssl_part, re.IGNORECASE)
            if m:
                ssl_val = m.group(1).lower()
        result["probe_ssl"] = ssl_val
        
        # 計算網速
        net_bytes = parse_net_output(net_part)
        current_time = time.time()
        
        rx_speed_str = "0.0 B/s"
        tx_speed_str = "0.0 B/s"
        
        if server_id in LAST_NET_STATS:
            prev = LAST_NET_STATS[server_id]
            dt = current_time - prev["time"]
            if dt > 0.05: # 防止除以零
                rx_speed = (net_bytes["rx"] - prev["rx"]) / dt
                tx_speed = (net_bytes["tx"] - prev["tx"]) / dt
                if rx_speed < 0: rx_speed = 0
                if tx_speed < 0: tx_speed = 0
                rx_speed_str = format_speed(rx_speed)
                tx_speed_str = format_speed(tx_speed)
                
        # 更新快取
        LAST_NET_STATS[server_id] = {
            "rx": net_bytes["rx"],
            "tx": net_bytes["tx"],
            "time": current_time
        }
        
        result["network"] = {
            "rx_speed": rx_speed_str,
            "tx_speed": tx_speed_str
        }
        
        return result
    finally:
        ssh.close()

def check_server_connection_quick(s):
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        port = int(s.get("port", 22))
        pwd = decrypt_password(s.get("password", ""))
        key_path = s.get("key_path")
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(s.get("host"), port=port, username=s.get("username"), pkey=key, timeout=1.5, disabled_algorithms=None)
        else:
            ssh.connect(s.get("host"), port=port, username=s.get("username"), password=pwd, timeout=1.5, disabled_algorithms=None)
        ssh.close()
        return True
    except Exception:
        return False

def fetch_windows_data(server):
    host = server.get("host")
    port = int(server.get("port", 5985))
    if port == 22:
        port = 5985  # 自動將預設 SSH 埠轉換為 WinRM 預設埠
    username = server.get("username")
    password = server.get("password")
    
    try:
        url = f"http://{host}:{port}/wsman"
        session = winrm.Session(url, auth=(username, password), transport='ntlm')
        
        # 取得所有正式 Debian Probe IP 清單（排除 Windows 主機與 Probe 測試設備）
        servers_list = read_servers()
        debian_ips = [
            s.get("host") for s in servers_list 
            if s.get("os_type", "debian") != "windows" 
            and s.get("category") != "__probe_test__" 
            and s.get("host")
        ]
        total = len(debian_ips)
        
        # PowerShell 指令：讀取 CoreServiceLog4Net.cfg 中的 LOG 等級 (使用安全的單引號語法)
        ps_log_test = "$logLevel = '未知'; $cfgPath = 'C:\\PIXIS\\CoreService\\Config\\CoreServiceLog4Net.cfg'; if (Test-Path $cfgPath) { try { [xml]$xml = Get-Content $cfgPath -Encoding UTF8; $node = $xml.SelectSingleNode('//root/level'); if ($node) { $logLevel = $node.getAttribute('value') } } catch {} }; Write-Output \"LOG_LEVEL:$logLevel\""
        # PowerShell 指令：讀取 CoreServiceSetting.cfg 中的 EnableSSL 設定
        ps_ssl_test = "$sslVal = '未知'; $c1 = 'C:\\PIXIS\\CoreService\\Config\\CoreServiceSetting.cfg'; if (-not (Test-Path $c1)) { $c1 = 'C:\\PIXIS\\CoreService\\Config\\CoreServiceSetting' }; if (Test-Path $c1) { try { $raw = Get-Content $c1 -Raw -Encoding UTF8; if ($raw -match '\"EnableSSL\"\\s*:\\s*(true|false)') { $sslVal = $matches[1].ToLower() } } catch {} }; Write-Output \"SERVER_SSL:$sslVal\""

        if total > 0:
            ips_array_str = ", ".join([f'"{ip}"' for ip in debian_ips])
            # 遠端 PowerShell：使用 TcpClient 在 300 毫秒內測試 SSH (22) 連接埠
            ps_conn_test = f'$ips = @({ips_array_str}); $connected = 0; foreach ($ip in $ips) {{ $t = New-Object System.Net.Sockets.TcpClient; $ok = $false; try {{ $res = $t.ConnectAsync($ip, 22).Wait(300); if ($t.Connected) {{ $ok = $true; $t.Close() }} }} catch {{}}; if ($ok) {{ $connected++ }} }}; Write-Output "CONNECTED_PROBES:$connected"'
            cmd = f'$cpuObj = Get-CimInstance -ClassName Win32_Processor | Measure-Object -Property LoadPercentage -Average; $cpu = if ($cpuObj.Average -ne $null) {{ [Math]::Round($cpuObj.Average) }} else {{ 0 }}; $os = Get-CimInstance -ClassName Win32_OperatingSystem; $memTotal = $os.TotalVisibleMemorySize; $memFree = $os.FreePhysicalMemory; $disk = Get-CimInstance -ClassName Win32_LogicalDisk -Filter \'DeviceID="C:"\' -ErrorAction SilentlyContinue; if (!$disk) {{ $disk = Get-CimInstance -ClassName Win32_LogicalDisk | Where-Object {{ $_.DeviceID -like "*C:*" }} }}; $diskSize = $disk.Size; $diskFree = $disk.FreeSpace; {ps_conn_test}; {ps_log_test}; {ps_ssl_test}; Write-Output "CPU:$cpu|MEM_TOTAL:$memTotal|MEM_FREE:$memFree|DISK_SIZE:$diskSize|DISK_FREE:$diskFree"'
        else:
            cmd = f'$cpuObj = Get-CimInstance -ClassName Win32_Processor | Measure-Object -Property LoadPercentage -Average; $cpu = if ($cpuObj.Average -ne $null) {{ [Math]::Round($cpuObj.Average) }} else {{ 0 }}; $os = Get-CimInstance -ClassName Win32_OperatingSystem; $memTotal = $os.TotalVisibleMemorySize; $memFree = $os.FreePhysicalMemory; $disk = Get-CimInstance -ClassName Win32_LogicalDisk -Filter \'DeviceID="C:"\' -ErrorAction SilentlyContinue; if (!$disk) {{ $disk = Get-CimInstance -ClassName Win32_LogicalDisk | Where-Object {{ $_.DeviceID -like "*C:*" }} }}; $diskSize = $disk.Size; $diskFree = $disk.FreeSpace; Write-Output "CONNECTED_PROBES:0"; {ps_log_test}; {ps_ssl_test}; Write-Output "CPU:$cpu|MEM_TOTAL:$memTotal|MEM_FREE:$memFree|DISK_SIZE:$diskSize|DISK_FREE:$diskFree"'
            
        r = session.run_ps(cmd)
        
        if r.status_code != 0:
            err_msg = r.std_err.decode('utf-8', errors='ignore')
            raise Exception(f"WinRM 遠端指令執行失敗 (Code {r.status_code}): {err_msg}")
            
        output = r.std_out.decode('utf-8', errors='ignore').strip()
            
        match = re.search(r'CPU:(\d+)\|MEM_TOTAL:(\d+)\|MEM_FREE:(\d+)\|DISK_SIZE:(\d+)\|DISK_FREE:(\d+)', output)
        if not match:
            raise Exception(f"解析失敗: {output}")
            
        cpu_val = float(match.group(1))
        mem_total_kb = float(match.group(2))
        mem_free_kb = float(match.group(3))
        disk_size_b = float(match.group(4))
        disk_free_b = float(match.group(5))
        
        mem_used_kb = mem_total_kb - mem_free_kb
        mem_percent = (mem_used_kb / mem_total_kb) * 100 if mem_total_kb > 0 else 0.0
        disk_used_b = disk_size_b - disk_free_b
        disk_percent = int((disk_used_b / disk_size_b) * 100) if disk_size_b > 0 else 0
        
        # 從遠端輸出中解析 Probe 連線數
        connected = 0
        connected_match = re.search(r'CONNECTED_PROBES:(\d+)', output)
        if connected_match:
            connected = int(connected_match.group(1))
                
        # 解析 LOG4NET 設定等級
        log_level = "未知"
        log_match = re.search(r'LOG_LEVEL:(.+)', output)
        if log_match:
            log_level = log_match.group(1).strip()

        # 解析 Server SSL (CoreServiceSetting.cfg)
        server_ssl = "未知"
        ssl_match = re.search(r'SERVER_SSL:(.+)', output)
        if ssl_match:
            server_ssl = ssl_match.group(1).strip()

        result = {
            "uptime": "Windows Server",
            "load_avg": [cpu_val / 100.0, cpu_val / 100.0, cpu_val / 100.0],
            "tasks": {"total": 0, "running": 0, "sleeping": 0, "stopped": 0, "zombie": 0},
            "cpu": {"user": cpu_val, "system": 0.0, "idle": 100.0 - cpu_val, "wait": 0.0},
            "mem": {
                "total": mem_total_kb / 1024.0, # MiB
                "used": mem_used_kb / 1024.0,
                "free": mem_free_kb / 1024.0,
                "cache": 0.0,
                "unit": "MiB"
            },
            "swap": {"total": 0.0, "used": 0.0, "free": 0.0, "avail": 0.0, "unit": "MiB"},
            "disk": {
                "total": f"{disk_size_b / (1024*1024*1024):.1f}G",
                "used": f"{disk_used_b / (1024*1024*1024):.1f}G",
                "avail": f"{disk_free_b / (1024*1024*1024):.1f}G",
                "percent": disk_percent
            },
            "network": {
                "rx_speed": "0.0 B/s",
                "tx_speed": "0.0 B/s"
            },
            "probe_status": {
                "connected": connected,
                "total": total,
                "detail": f"{connected}/{total}"
            },
            "log_level": log_level,
            "server_ssl": server_ssl,
            "processes": []
        }
        return result
    except Exception as e:
        raise e

# API 路由
@app.route('/')
def home():
    return render_template('index.html')

@app.route('/api/servers', methods=['GET'])
def get_servers():
    return jsonify(read_servers())

@app.route('/api/servers', methods=['POST'])
def add_server():
    data = request.json
    servers = read_servers()
    new_server = {
        "id": str(uuid.uuid4()),
        "name": data.get("name"),
        "category": data.get("category", "預設分類"),
        "host": data.get("host"),
        "port": int(data.get("port", 22)),
        "username": data.get("username"),
        "password": data.get("password", ""),
        "key_path": data.get("key_path", ""),
        "os_type": data.get("os_type", "debian")
    }
    servers.append(new_server)
    if write_servers(servers):
        return jsonify({"success": True, "server": new_server})
    return jsonify({"success": False, "message": "儲存伺服器設定時發生錯誤"}), 500

# NAC API 設定取得與儲存
@app.route('/api/settings/nac', methods=['GET'])
def get_nac_settings():
    server_ip = get_db_setting("nac_server_ip", "https://192.168.28.181:8001")
    account = get_db_setting("nac_account", "admin")
    enc_pwd = get_db_setting("nac_password", "")
    password = decrypt_password(enc_pwd) if enc_pwd else "admin"
    return jsonify({
        "nac_server_ip": server_ip,
        "nac_account": account,
        "nac_password": password
    })

@app.route('/api/settings/nac', methods=['POST'])
def save_nac_settings():
    data = request.json or {}
    server_ip = data.get("nac_server_ip", "").strip()
    account = data.get("nac_account", "").strip()
    password = data.get("nac_password", "")
    
    if not server_ip:
        return jsonify({"success": False, "message": "NAC Server IP 不能為空"}), 400
        
    set_db_setting("nac_server_ip", server_ip)
    set_db_setting("nac_account", account)
    if password is not None:
        set_db_setting("nac_password", encrypt_password(password))
        
    return jsonify({"success": True, "message": "NAC 設定已成功儲存"})

def get_athenac_token(server_ip, account, password):
    url = f"{server_ip.rstrip('/')}/api/connect/token"
    headers = {'Content-Type': 'application/x-www-form-urlencoded'}
    encoded_pwd = base64.b64encode(password.encode("utf-8")).decode("utf-8")
    form_data = {
        "grant_type": "password",
        "scope": "offline_access",
        "username": account,
        "password": encoded_pwd
    }
    response = requests.post(url, headers=headers, data=form_data, verify=False, timeout=10)
    if response.status_code == 200:
        res = response.json()
        return "Bearer " + res["access_token"]
    else:
        raise Exception(f"NAC 登入失敗 [{response.status_code}]: {response.text}")

@app.route('/api/probes/fetch-ips', methods=['POST'])
def fetch_probe_ips():
    data = request.json or {}
    server_ip = data.get("nac_server_ip") or get_db_setting("nac_server_ip", "https://192.168.28.181:8001")
    account = data.get("nac_account") or get_db_setting("nac_account", "admin")
    enc_pwd = get_db_setting("nac_password", "")
    password = data.get("nac_password") or (decrypt_password(enc_pwd) if enc_pwd else "admin")

    try:
        import urllib3
        urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

        # 1. 登入取得 Token
        token = get_athenac_token(server_ip, account, password)
        
        # 2. 打 /api/Probes (優先使用 POST，若回傳 405 則自動嘗試 GET)
        headers = {'Authorization': token, 'Content-Type': 'application/json'}
        probes_url = f"{server_ip.rstrip('/')}/api/Probes"
        resp = requests.post(probes_url, headers=headers, json={}, verify=False, timeout=10)
        if resp.status_code == 405:
            resp = requests.get(probes_url, headers={'Authorization': token}, verify=False, timeout=10)
        
        if resp.status_code == 200:
            res_data = resp.json()
            probes = res_data.get('Data', [])
            ip_set = set()
            for probe in probes:
                # 1. 主 Probe IP
                main_probe = probe.get('MainProbe', {})
                if isinstance(main_probe, dict):
                    main_ip = main_probe.get('IP')
                    if main_ip:
                        ip_set.add(str(main_ip).strip())
                # 2. 備援 Probe IP (如果有)
                standby_probe = probe.get('StandbyProbe')
                if isinstance(standby_probe, dict):
                    standby_ip = standby_probe.get('IP')
                    if standby_ip:
                        ip_set.add(str(standby_ip).strip())
            return jsonify({
                "success": True,
                "ips": sorted(list(ip_set)),
                "count": len(ip_set),
                "server_ip": server_ip
            })
        else:
            return jsonify({"success": False, "message": f"取得 Probes API 失敗 [{resp.status_code}]: {resp.text}"}), 400
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500

@app.route('/api/servers/batch', methods=['POST'])
def batch_add_servers():
    data = request.json or {}
    server_list = data.get("servers", [])
    category = data.get("category", "預設分類")
    
    if not server_list or not isinstance(server_list, list):
        return jsonify({"success": False, "message": "請提供要匯入的主機清單"}), 400

    servers = read_servers()
    added_count = 0

    for item in server_list:
        host_ip = item.get("host", "").strip()
        if not host_ip:
            continue
        username = item.get("username", "root").strip() or "root"
        password = item.get("password", "1111")
        name = item.get("name", host_ip).strip() or host_ip
        
        item_cat = item.get("category", "").strip() or category
        if item_cat:
            add_custom_category_db("debian", item_cat)
        
        new_server = {
            "id": str(uuid.uuid4()),
            "name": name,
            "category": item_cat,
            "host": host_ip,
            "port": int(item.get("port", 22)),
            "username": username,
            "password": password,
            "key_path": "",
            "os_type": "debian"
        }
        servers.append(new_server)
        added_count += 1

    if write_servers(servers):
        return jsonify({"success": True, "message": f"成功匯入 {added_count} 台主機"})
    return jsonify({"success": False, "message": "儲存至資料庫時發生錯誤"}), 500

@app.route('/api/servers/parse-excel-csv', methods=['POST'])
def parse_excel_csv():
    import csv
    if 'file' not in request.files:
        return jsonify({"success": False, "message": "未收到上傳的檔案"}), 400
    
    file = request.files['file']
    if not file or file.filename == '':
        return jsonify({"success": False, "message": "未選取任何檔案"}), 400
        
    filename = file.filename.lower()
    raw_rows = []
    
    # 讀取檔案內容為文字或單元格資料
    if filename.endswith('.csv'):
        content = file.read()
        decoded_text = None
        for enc in ['utf-8-sig', 'utf-8', 'cp950', 'big5', 'gbk', 'latin1']:
            try:
                decoded_text = content.decode(enc)
                break
            except Exception:
                continue
        if decoded_text is None:
            return jsonify({"success": False, "message": "無法解析 CSV 編碼，請使用 UTF-8 或 Big5"}), 400
        
        reader = csv.reader(io.StringIO(decoded_text))
        for r in reader:
            raw_rows.append([str(c).strip() for c in r])
            
    elif filename.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if any(row):
                    raw_rows.append([str(c).strip() if c is not None else "" for c in row])
        except Exception as e:
            return jsonify({"success": False, "message": f"解析 Excel 失敗: {str(e)}"}), 400
    elif filename.endswith('.xls'):
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=file.read())
            sheet = book.sheet_by_index(0)
            for rx in range(sheet.nrows):
                row = sheet.row_values(rx)
                raw_rows.append([str(c).strip() for c in row])
        except ImportError:
            return jsonify({"success": False, "message": "伺服器未安裝 xlrd 套件以支援舊版 .xls 格式，請轉存為 .xlsx 或 .csv"}), 400
        except Exception as e:
            return jsonify({"success": False, "message": f"解析 .xls 失敗: {str(e)}"}), 400
    else:
        return jsonify({"success": False, "message": "不支援的檔案格式，請上傳 .xlsx, .xls 或 .csv"}), 400

    if not raw_rows:
        return jsonify({"success": False, "message": "檔案中沒有可用的資料行"}), 400

    # 智慧欄位辨識關鍵字
    ip_keywords = ['probe ip', 'probe_ip', 'probeip', 'probe', 'host', 'ip', 'ip位址', '主機ip', '伺服器ip', '設備ip', 'ip address']
    user_keywords = ['username', 'user', 'account', 'ssh帳號', '帳號', '使用者', '登入帳號', '用戶名']
    pwd_keywords = ['password', 'pass', 'pwd', 'ssh密碼', '密碼', '登入密碼']
    cat_keywords = ['category', 'cat', '分類', '群組', '目錄', '分類目錄']
    name_keywords = ['name', 'hostname', '主機名稱', '設備名稱', '名稱']

    ip_regex = re.compile(r'^(\d{1,3}\.){3}\d{1,3}$')

    def is_ip(val):
        val = str(val).strip()
        if not ip_regex.match(val):
            return False
        parts = val.split('.')
        return all(0 <= int(p) <= 255 for p in parts)

    # 尋找表頭列（在前 10 列中尋找是否有包含關鍵字的列）
    header_row_idx = -1
    col_mapping = {'ip': -1, 'user': -1, 'pwd': -1, 'cat': -1, 'name': -1}

    for idx, row in enumerate(raw_rows[:10]):
        row_lower = [c.lower() for c in row]
        ip_col = next((i for i, c in enumerate(row_lower) if any(k == c or k in c for k in ip_keywords)), -1)
        if ip_col != -1:
            header_row_idx = idx
            col_mapping['ip'] = ip_col
            col_mapping['user'] = next((i for i, c in enumerate(row_lower) if any(k == c or k in c for k in user_keywords)), -1)
            col_mapping['pwd'] = next((i for i, c in enumerate(row_lower) if any(k == c or k in c for k in pwd_keywords)), -1)
            col_mapping['cat'] = next((i for i, c in enumerate(row_lower) if any(k == c or k in c for k in cat_keywords)), -1)
            col_mapping['name'] = next((i for i, c in enumerate(row_lower) if i != ip_col and any(k == c or k in c for k in name_keywords)), -1)
            break

    parsed_servers = []
    seen_ips = set()

    data_rows = raw_rows[header_row_idx + 1:] if header_row_idx != -1 else raw_rows

    for row in data_rows:
        if not any(row):
            continue

        ip_val = ""
        user_val = ""
        pwd_val = ""
        cat_val = ""
        name_val = ""

        # 優先從識別出的 IP 欄位取出
        if col_mapping['ip'] != -1 and col_mapping['ip'] < len(row):
            candidate = row[col_mapping['ip']].strip()
            ip_match = re.search(r'(\d{1,3}\.){3}\d{1,3}', candidate)
            if ip_match and is_ip(ip_match.group(0)):
                ip_val = ip_match.group(0)

        # 若未透過表頭找到 IP，則自動掃描該列所有儲存格尋找 IP
        if not ip_val:
            for c_idx, cell in enumerate(row):
                c_str = str(cell).strip()
                ip_match = re.search(r'(\d{1,3}\.){3}\d{1,3}', c_str)
                if ip_match and is_ip(ip_match.group(0)):
                    ip_val = ip_match.group(0)
                    break

        if not ip_val or ip_val in seen_ips:
            continue

        seen_ips.add(ip_val)

        if col_mapping['user'] != -1 and col_mapping['user'] < len(row):
            user_val = row[col_mapping['user']].strip()
        if col_mapping['pwd'] != -1 and col_mapping['pwd'] < len(row):
            pwd_val = row[col_mapping['pwd']].strip()
        if col_mapping['cat'] != -1 and col_mapping['cat'] < len(row):
            cat_val = row[col_mapping['cat']].strip()
        if col_mapping['name'] != -1 and col_mapping['name'] < len(row):
            name_val = row[col_mapping['name']].strip()

        parsed_servers.append({
            "ip": ip_val,
            "name": name_val or ip_val,
            "username": user_val or "root",
            "password": pwd_val or "1111",
            "category": cat_val or ""
        })

    if not parsed_servers:
        return jsonify({"success": False, "message": "未能從檔案中解析出任何有效的 Probe IP"}), 400

    return jsonify({
        "success": True,
        "message": f"成功從檔案解析出 {len(parsed_servers)} 筆主機資料",
        "servers": parsed_servers
    })

@app.route('/api/servers/batch-delete', methods=['POST'])
def batch_delete_servers():
    data = request.json or {}
    server_ids = data.get("server_ids", [])
    if not server_ids or not isinstance(server_ids, list):
        return jsonify({"success": False, "message": "請提供有效的刪除主機 ID 清單"}), 400

    servers = read_servers()
    new_servers = [s for s in servers if s["id"] not in server_ids]
    deleted_count = len(servers) - len(new_servers)

    if write_servers(new_servers):
        return jsonify({"success": True, "message": f"成功刪除 {deleted_count} 台主機"})
    return jsonify({"success": False, "message": "儲存至資料庫時發生錯誤"}), 500

@app.route('/api/servers/<server_id>', methods=['PUT'])
def update_server(server_id):
    data = request.json
    servers = read_servers()
    for idx, s in enumerate(servers):
        if s["id"] == server_id:
            servers[idx].update({
                "name": data.get("name", s["name"]),
                "category": data.get("category", s["category"]),
                "host": data.get("host", s["host"]),
                "port": int(data.get("port", s["port"])),
                "username": data.get("username", s["username"]),
                "password": data.get("password", s["password"]),
                "key_path": data.get("key_path", s["key_path"]),
                "os_type": data.get("os_type", s.get("os_type", "debian"))
            })
            if write_servers(servers):
                return jsonify({"success": True, "server": servers[idx]})
            return jsonify({"success": False, "message": "儲存伺服器設定時發生錯誤"}), 500
    return jsonify({"success": False, "message": "找不到該伺服器"}), 404

@app.route('/api/servers/<server_id>', methods=['DELETE'])
def delete_server(server_id):
    servers = read_servers()
    new_servers = [s for s in servers if s["id"] != server_id]
    if len(new_servers) == len(servers):
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
    if write_servers(new_servers):
        return jsonify({"success": True})
    return jsonify({"success": False, "message": "儲存伺服器設定時發生錯誤"}), 500

def get_custom_categories(os_type):
    key = f"custom_categories_{os_type}"
    raw = get_db_setting(key, "[]")
    try:
        cats = json.loads(raw)
        return cats if isinstance(cats, list) else []
    except Exception:
        return []

def add_custom_category_db(os_type, category_name):
    cats = get_custom_categories(os_type)
    if category_name not in cats:
        cats.append(category_name)
        set_db_setting(f"custom_categories_{os_type}", json.dumps(cats))

def remove_custom_category_db(os_type, category_name):
    cats = get_custom_categories(os_type)
    if category_name in cats:
        cats = [c for c in cats if c != category_name]
        set_db_setting(f"custom_categories_{os_type}", json.dumps(cats))

def rename_custom_category_db(os_type, old_name, new_name):
    cats = get_custom_categories(os_type)
    if old_name in cats:
        cats = [new_name if c == old_name else c for c in cats]
        set_db_setting(f"custom_categories_{os_type}", json.dumps(cats))

@app.route('/api/categories', methods=['GET'])
def get_categories():
    servers = read_servers()
    win_cats = [s.get("category") for s in servers if s.get("os_type") == "windows" and s.get("category")]
    deb_cats = [s.get("category") for s in servers if s.get("os_type") != "windows" and s.get("category") and s.get("category") != "__probe_test__"]
    
    custom_win = get_custom_categories("windows")
    custom_deb = get_custom_categories("debian")
    
    all_win = sorted(list(set(win_cats + custom_win)))
    all_deb = sorted(list(set(deb_cats + custom_deb)))
    
    if "預設分類" in all_win: all_win.remove("預設分類")
    all_win.insert(0, "預設分類")
    
    if "預設分類" in all_deb: all_deb.remove("預設分類")
    all_deb.insert(0, "預設分類")
    
    return jsonify({
        "windows": all_win,
        "debian": all_deb
    })

@app.route('/api/categories/add', methods=['POST'])
def add_category():
    data = request.json or {}
    cat_name = data.get("category", "").strip()
    os_type = data.get("os_type", "debian")
    
    if not cat_name:
        return jsonify({"success": False, "message": "群組名稱不能為空"}), 400
        
    add_custom_category_db(os_type, cat_name)
    return jsonify({"success": True, "message": "群組新增成功"})

@app.route('/api/servers/rename-category', methods=['POST'])
def rename_category():
    data = request.json
    old_cat = data.get("old_category")
    new_cat = data.get("new_category", "").strip()
    os_type = data.get("os_type", "debian")
    
    if not old_cat or not new_cat:
        return jsonify({"success": False, "message": "舊群組名稱與新群組名稱皆不可為空"}), 400
        
    try:
        rename_custom_category_db(os_type, old_cat, new_cat)
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE servers SET category = ? WHERE category = ? AND (os_type = ? OR (? = 'debian' AND os_type IS NULL))",
            (new_cat, old_cat, os_type, os_type)
        )
        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "message": f"批次修改群組失敗: {str(e)}"}), 500

@app.route('/api/servers/delete-category', methods=['POST'])
def delete_category():
    data = request.json
    cat = data.get("category")
    os_type = data.get("os_type", "debian")
    action = data.get("action", "move_to_default") # move_to_default 或 delete_servers
    
    if not cat:
        return jsonify({"success": False, "message": "群組名稱不可為空"}), 400
        
    try:
        remove_custom_category_db(os_type, cat)
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        if action == "delete_servers":
            cursor.execute(
                "DELETE FROM servers WHERE category = ? AND (os_type = ? OR (? = 'debian' AND os_type IS NULL))",
                (cat, os_type, os_type)
            )
        else: # move_to_default
            cursor.execute(
                "UPDATE servers SET category = '預設分類' WHERE category = ? AND (os_type = ? OR (? = 'debian' AND os_type IS NULL))",
                (cat, os_type, os_type)
            )
        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "message": f"刪除群組失敗: {str(e)}"}), 500

@app.route('/api/servers/<server_id>/top', methods=['POST'])
def get_server_top(server_id):
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
    
    os_type = server.get("os_type", "debian")
    attempts = 2
    last_err = None
    for attempt in range(attempts):
        try:
            if os_type == "windows":
                top_data = fetch_windows_data(server)
            else:
                top_data = fetch_top_data(server)
            return jsonify({"success": True, "data": top_data})
        except Exception as e:
            last_err = str(e)
            if attempt < attempts - 1:
                time.sleep(0.1) # 輕微延遲後立即重試一次
                
    return jsonify({"success": False, "message": last_err}), 500

# 測試連線 API
@app.route('/api/servers/test-connect', methods=['POST'])
def test_connect():
    data = request.json
    os_type = data.get("os_type", "debian")
    try:
        if os_type == "windows":
            fetch_windows_data(data)
        else:
            fetch_top_data(data)
        return jsonify({"success": True, "message": "連線成功！"})
    except Exception as e:
        return jsonify({"success": False, "message": f"連線失敗: {str(e)}"})

def read_templates():
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM maintenance_templates")
        rows = cursor.fetchall()
        templates = [dict(row) for row in rows]
        conn.close()
        return templates
    except Exception as e:
        print(f"Error reading templates from DB: {e}")
        return []

def write_templates(templates):
    try:
        conn = sqlite3.connect(DB_FILE_PATH)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM maintenance_templates")
        for t in templates:
            cursor.execute('''
                INSERT INTO maintenance_templates (id, name, customer, contact, version, location)
                VALUES (?, ?, ?, ?, ?, ?)
            ''', (
                t.get("id"),
                t.get("name", "未命名範本"),
                t.get("customer", ""),
                t.get("contact", ""),
                t.get("version", ""),
                t.get("location", "")
            ))
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print(f"Error writing templates to DB: {e}")
        return False

@app.route('/api/maintenance-templates', methods=['GET'])
def get_templates():
    return jsonify(read_templates())

@app.route('/api/maintenance-templates', methods=['POST'])
def add_template():
    data = request.json
    templates = read_templates()
    new_template = {
        "id": str(uuid.uuid4()),
        "name": data.get("name", "未命名範本"),
        "customer": data.get("customer", ""),
        "contact": data.get("contact", ""),
        "version": data.get("version", ""),
        "location": data.get("location", "")
    }
    templates.append(new_template)
    if write_templates(templates):
        return jsonify({"success": True, "template": new_template})
    return jsonify({"success": False, "message": "儲存範本時發生錯誤"}), 500

@app.route('/api/maintenance-templates/<template_id>', methods=['PUT'])
def update_template(template_id):
    data = request.json
    templates = read_templates()
    for idx, t in enumerate(templates):
        if t["id"] == template_id:
            templates[idx].update({
                "name": data.get("name", t["name"]),
                "customer": data.get("customer", t["customer"]),
                "contact": data.get("contact", t["contact"]),
                "version": data.get("version", t["version"]),
                "location": data.get("location", t["location"])
            })
            if write_templates(templates):
                return jsonify({"success": True, "template": templates[idx]})
            return jsonify({"success": False, "message": "儲存範本時發生錯誤"}), 500
    return jsonify({"success": False, "message": "找不到該範本"}), 404

@app.route('/api/maintenance-templates/<template_id>', methods=['DELETE'])
def delete_template(template_id):
    templates = read_templates()
    new_templates = [t for t in templates if t["id"] != template_id]
    if len(new_templates) == len(templates):
        return jsonify({"success": False, "message": "找不到該範本"}), 404
    if write_templates(new_templates):
        return jsonify({"success": True})
    return jsonify({"success": False, "message": "儲存範本時發生錯誤"}), 500

@app.route('/api/servers/<server_id>/export-excel', methods=['POST'])
def export_excel(server_id):
    data = request.json
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
        
    os_type = server.get("os_type", "debian")
    ip = server.get("host")
    cpu_usage = data.get("cpu_usage", "0")
    mem_usage = data.get("mem_usage", "0")
    disk_avail = data.get("disk_avail", "0G")
    probe_detail = data.get("probe_detail", "0/0")
    
    customer = data.get("customer", "")
    contact = data.get("contact", "")
    version = data.get("version", "")
    location = data.get("location", "")
    
    if not ip:
        return jsonify({"success": False, "message": "缺少主機 IP 資訊"}), 400
        
    try:
        if os_type == "windows":
            # 載入 Windows Server 範本並修改
            template_path = os.path.join(get_base_dir(), "PIXIS 維護表單_Server.xlsx")
            if not os.path.exists(template_path):
                return jsonify({"success": False, "message": "找不到維護單範本 PIXIS 維護表單_Server.xlsx"}), 404
                
            wb = openpyxl.load_workbook(template_path)
            sheet = wb.active
            
            sheet['D4'].value = f"維護日期：{datetime.now().strftime('%Y-%m-%d')}"
            if customer: sheet['B4'].value = customer
            if contact: sheet['B5'].value = contact
            if version: sheet['D5'].value = f"PIXIS 版本：{version}"
            if location: sheet['D6'].value = f"地點：{location}"
            
            sheet['D9'].value = "○"
            sheet['E9'].value = f"本次檢查中，常態維持 {cpu_usage}% 以下"
            
            sheet['D10'].value = "○"
            sheet['E10'].value = f"本次檢查中，常態維持 {mem_usage}% 以下"
            
            sheet['D11'].value = "○"
            sheet['E11'].value = f"C槽剩餘 {disk_avail}"
            
            sheet['D12'].value = ip
            
            sheet['D13'].value = "○"
            sheet['E13'].value = "Service ON，服務正常"
            
            sheet['D14'].value = "○"
            sheet['E14'].value = f"Probe 連線數 {probe_detail}"
            
            timestamp = datetime.now().strftime("%Y%m%d%H%M")
            file_stream = io.BytesIO()
            wb.save(file_stream)
            file_stream.seek(0)
            
            return send_file(
                file_stream,
                as_attachment=True,
                download_name=f"{ip}_Server維護單_{timestamp}.xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            # 載入 Debian Probe 範本並修改
            template_path = os.path.join(get_base_dir(), "PIXIS 維護表單_Probe.xlsx")
            if not os.path.exists(template_path):
                return jsonify({"success": False, "message": "找不到維護單範本 PIXIS 維護表單_Probe.xlsx"}), 404
                
            wb = openpyxl.load_workbook(template_path)
            sheet = wb.active
            
            sheet['D4'].value = f"維護日期：{datetime.now().strftime('%Y-%m-%d')}"
            if customer: sheet['B4'].value = customer
            if contact: sheet['B5'].value = contact
            if version: sheet['D5'].value = f"PIXIS 版本：{version}"
            if location: sheet['D6'].value = f"地點：{location}"
            
            sheet['D9'].value = ip
            
            sheet['D11'].value = "○"
            sheet['E11'].value = f"本次檢查中，常態維持{cpu_usage}% 以下"
            
            sheet['D12'].value = "○"
            sheet['E12'].value = f"本次檢查中，常態維持 {mem_usage}% 以下"
            
            timestamp = datetime.now().strftime("%Y%m%d%H%M")
            file_stream = io.BytesIO()
            wb.save(file_stream)
            file_stream.seek(0)
            
            return send_file(
                file_stream,
                as_attachment=True,
                download_name=f"{ip}_Probe維護單_{timestamp}.xlsx",
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
    except Exception as e:
        return jsonify({"success": False, "message": f"匯出失敗: {str(e)}"}), 500

@app.route('/api/servers/batch-export-excel', methods=['POST'])
def batch_export_excel():
    data = request.json or {}
    server_ids = data.get("server_ids", [])
    if not server_ids:
        return jsonify({"success": False, "message": "請至少選擇一台目標設備"}), 400

    customer = data.get("customer", "")
    contact = data.get("contact", "")
    version = data.get("version", "")
    location = data.get("location", "")

    all_servers = read_servers()
    target_servers = [s for s in all_servers if s["id"] in server_ids]

    if not target_servers:
        return jsonify({"success": False, "message": "找不到指定的設備"}), 404

    # 平行抓取所選設備的即時數據
    def fetch_server_metrics(server):
        os_type = server.get("os_type", "debian")
        try:
            if os_type == "windows":
                top_data = fetch_windows_data(server)
                used_cpu = round(100 - top_data.get("cpu", {}).get("idle", 100), 1)
                mem = top_data.get("mem", {})
                mem_total = mem.get("total", 1)
                mem_used = mem.get("used", 0)
                used_mem = round((mem_used / mem_total) * 100, 1) if mem_total else 0
                disk_avail = top_data.get("disk", {}).get("avail", "0G")
                probe_detail = top_data.get("probe_status", {}).get("detail", "0/0")
                return {
                    "cpu_usage": str(used_cpu),
                    "mem_usage": str(used_mem),
                    "disk_avail": str(disk_avail),
                    "probe_detail": str(probe_detail)
                }
            else:
                top_data = fetch_top_data(server)
                used_cpu = round(100 - top_data.get("cpu", {}).get("idle", 100), 1)
                mem = top_data.get("mem", {})
                mem_total = mem.get("total", 1)
                mem_used = mem.get("used", 0)
                used_mem = round((mem_used / mem_total) * 100, 1) if mem_total else 0
                return {
                    "cpu_usage": str(used_cpu),
                    "mem_usage": str(used_mem),
                    "disk_avail": "0G",
                    "probe_detail": "0/0"
                }
        except Exception as e:
            print(f"抓取設備 {server.get('host')} 數據失敗: {e}")
            return {
                "cpu_usage": "0",
                "mem_usage": "0",
                "disk_avail": "0G",
                "probe_detail": "0/0"
            }

    metrics_map = {}
    max_workers = min(10, len(target_servers))
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        future_map = {executor.submit(fetch_server_metrics, s): s["id"] for s in target_servers}
        for future in concurrent.futures.as_completed(future_map):
            s_id = future_map[future]
            try:
                metrics_map[s_id] = future.result()
            except Exception:
                metrics_map[s_id] = {"cpu_usage": "0", "mem_usage": "0", "disk_avail": "0G", "probe_detail": "0/0"}

    timestamp = datetime.now().strftime("%Y%m%d%H%M")
    zip_buffer = io.BytesIO()

    try:
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for server in target_servers:
                os_type = server.get("os_type", "debian")
                ip = server.get("host", "unknown")
                metrics = metrics_map.get(server["id"], {"cpu_usage": "0", "mem_usage": "0", "disk_avail": "0G", "probe_detail": "0/0"})

                cpu_usage = metrics["cpu_usage"]
                mem_usage = metrics["mem_usage"]
                disk_avail = metrics["disk_avail"]
                probe_detail = metrics["probe_detail"]

                if os_type == "windows":
                    template_path = os.path.join(get_base_dir(), "PIXIS 維護表單_Server.xlsx")
                    file_suffix = "Server維護單"
                else:
                    template_path = os.path.join(get_base_dir(), "PIXIS 維護表單_Probe.xlsx")
                    file_suffix = "Probe維護單"

                if not os.path.exists(template_path):
                    continue

                wb = openpyxl.load_workbook(template_path)
                sheet = wb.active

                sheet['D4'].value = f"維護日期：{datetime.now().strftime('%Y-%m-%d')}"
                if customer: sheet['B4'].value = customer
                if contact: sheet['B5'].value = contact
                if version: sheet['D5'].value = f"PIXIS 版本：{version}"
                if location: sheet['D6'].value = f"地點：{location}"

                if os_type == "windows":
                    sheet['D9'].value = "○"
                    sheet['E9'].value = f"本次檢查中，常態維持 {cpu_usage}% 以下"

                    sheet['D10'].value = "○"
                    sheet['E10'].value = f"本次檢查中，常態維持 {mem_usage}% 以下"

                    sheet['D11'].value = "○"
                    sheet['E11'].value = f"C槽剩餘 {disk_avail}"

                    sheet['D12'].value = ip

                    sheet['D13'].value = "○"
                    sheet['E13'].value = "Service ON，服務正常"

                    sheet['D14'].value = "○"
                    sheet['E14'].value = f"Probe 連線數 {probe_detail}"
                else:
                    sheet['D9'].value = ip

                    sheet['D11'].value = "○"
                    sheet['E11'].value = f"本次檢查中，常態維持{cpu_usage}% 以下"

                    sheet['D12'].value = "○"
                    sheet['E12'].value = f"本次檢查中，常態維持 {mem_usage}% 以下"

                file_stream = io.BytesIO()
                wb.save(file_stream)
                excel_bytes = file_stream.getvalue()

                filename = f"{ip}_{file_suffix}_{timestamp}.xlsx"
                zip_file.writestr(filename, excel_bytes)

        zip_buffer.seek(0)
        return send_file(
            zip_buffer,
            as_attachment=True,
            download_name=f"維護單批次打包_{timestamp}.zip",
            mimetype="application/zip"
        )
    except Exception as e:
        return jsonify({"success": False, "message": f"批次匯出維護單失敗: {str(e)}"}), 500

@app.route('/api/servers/<server_id>/log-level', methods=['POST'])
def update_server_log_level(server_id):
    data = request.json or {}
    new_level = data.get("log_level", "").upper().strip()
    if new_level not in ["DEBUG", "INFO", "WARN"]:
        return jsonify({"success": False, "message": "不支援的 LOG 等級，僅限 DEBUG / INFO / WARN"}), 400
        
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server or server.get("os_type") != "windows":
        return jsonify({"success": False, "message": "找不到該 Windows 伺服器"}), 404
        
    host = server.get("host")
    port = int(server.get("port", 5985))
    if port == 22: port = 5985
    username = server.get("username")
    password = server.get("password")
    
    try:
        url = f"http://{host}:{port}/wsman"
        session = winrm.Session(url, auth=(username, password), transport='ntlm')
        ps_script = f'''
        $cfgPath = 'C:\\PIXIS\\CoreService\\Config\\CoreServiceLog4Net.cfg';
        if (-not (Test-Path $cfgPath)) {{
            Write-Output "ERROR:找不到設定檔 $cfgPath";
            exit 1;
        }}
        try {{
            [xml]$xml = Get-Content $cfgPath -Encoding UTF8;
            $node = $xml.SelectSingleNode('//root/level');
            if ($node) {{
                $node.SetAttribute('value', '{new_level}');
                $xml.Save($cfgPath);
                Write-Output "SUCCESS";
            }} else {{
                Write-Output "ERROR: XML 缺少 <level> 標籤";
                exit 1;
            }}
        }} catch {{
            Write-Output "ERROR: " + $_.Exception.Message;
            exit 1;
        }}
        '''
        r = session.run_ps(ps_script)
        out = r.std_out.decode('utf-8', errors='ignore').strip()
        if "SUCCESS" in out:
            return jsonify({"success": True, "message": f"LOG 設定已成功更新為 {new_level}"})
        else:
            return jsonify({"success": False, "message": f"更新失敗: {out}"}), 500
    except Exception as e:
        return jsonify({"success": False, "message": f"WinRM 連線或執行異常: {str(e)}"}), 500

@app.route('/api/servers/<server_id>/server-ssl-files', methods=['GET'])
def get_server_ssl_files(server_id):
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server or server.get("os_type") != "windows":
        return jsonify({"success": False, "message": "找不到該 Windows 伺服器"}), 404

    host = server.get("host")
    port = int(server.get("port", 5985))
    if port == 22: port = 5985
    username = server.get("username")
    password = server.get("password")

    try:
        url = f"http://{host}:{port}/wsman"
        session = winrm.Session(url, auth=(username, password), transport='ntlm')
        ps_script = '''
        $files = @();
        $candidates = @(
            'C:\\PIXIS\\CoreService\\Config\\CoreServiceSetting.cfg',
            'C:\\PIXIS\\CoreService\\Config\\CoreServiceSetting',
            'C:\\PIXIS\\Web\\Config\\PixisWebSetting.cfg',
            'C:\\PIXIS\\Web\\Config\\PixisWebSetting'
        );
        foreach ($p in $candidates) {
            if (Test-Path $p) {
                $base = $p -replace '\\.cfg$', '';
                $already = $false;
                foreach ($f in $files) {
                    if (($f -replace '\\.cfg$', '') -eq $base) { $already = $true; break; }
                }
                if (-not $already) { $files += $p; }
            }
        }
        $files -join '|';
        '''
        r = session.run_ps(ps_script)
        out = r.std_out.decode('utf-8', errors='ignore').strip()
        files = [f for f in out.split('|') if f]
        return jsonify({"success": True, "files": files, "count": len(files)})
    except Exception as e:
        return jsonify({"success": False, "message": f"WinRM 連線或掃描失敗: {str(e)}"}), 500

@app.route('/api/servers/<server_id>/server-ssl', methods=['POST'])
def update_server_ssl(server_id):
    data = request.json or {}
    enable_ssl = data.get("enable_ssl")
    if enable_ssl is None:
        return jsonify({"success": False, "message": "缺少 enable_ssl 參數"}), 400

    target_bool_str = "true" if enable_ssl is True or str(enable_ssl).lower() == "true" else "false"

    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server or server.get("os_type") != "windows":
        return jsonify({"success": False, "message": "找不到該 Windows 伺服器"}), 404

    host = server.get("host")
    port = int(server.get("port", 5985))
    if port == 22: port = 5985
    username = server.get("username")
    password = server.get("password")

    try:
        url = f"http://{host}:{port}/wsman"
        session = winrm.Session(url, auth=(username, password), transport='ntlm')
        ps_script = f'''
        $target = "{target_bool_str}";
        $filesToUpdate = @(
            'C:\\PIXIS\\CoreService\\Config\\CoreServiceSetting.cfg',
            'C:\\PIXIS\\CoreService\\Config\\CoreServiceSetting',
            'C:\\PIXIS\\Web\\Config\\PixisWebSetting.cfg',
            'C:\\PIXIS\\Web\\Config\\PixisWebSetting'
        );
        $updated = @();
        $handledBases = @();
        foreach ($p in $filesToUpdate) {{
            $base = $p -replace '\\.cfg$', '';
            if ($handledBases -contains $base) {{ continue; }}
            if (Test-Path $p) {{
                try {{
                    $content = [System.IO.File]::ReadAllText($p, [System.Text.Encoding]::UTF8);
                    if ($content -match '("EnableSSL"\\s*:\\s*)(true|false)') {{
                        $newContent = [System.Text.RegularExpressions.Regex]::Replace($content, '("EnableSSL"\\s*:\\s*)(true|false)', "`$1$target", [System.Text.RegularExpressions.RegexOptions]::IgnoreCase);
                        [System.IO.File]::WriteAllText($p, $newContent, [System.Text.Encoding]::UTF8);
                        $updated += $p;
                        $handledBases += $base;
                    }}
                }} catch {{}}
            }}
        }}
        if ($updated.Count -gt 0) {{
            Write-Output "SUCCESS:" + ($updated -join '|');
        }} else {{
            Write-Output "ERROR:未找到或未更新任何設定檔";
        }}
        '''
        r = session.run_ps(ps_script)
        out = r.std_out.decode('utf-8', errors='ignore').strip()
        if "SUCCESS:" in out:
            files_str = out.split("SUCCESS:")[1].strip()
            files_list = files_str.split('|')
            return jsonify({
                "success": True,
                "message": f"SSL 通訊設定已成功變更為 {target_bool_str} (共更新 {len(files_list)} 個設定檔)",
                "modified": files_list
            })
        else:
            return jsonify({"success": False, "message": f"更新失敗: {out}"}), 500
    except Exception as e:
        return jsonify({"success": False, "message": f"WinRM 連線或執行異常: {str(e)}"}), 500

@app.route('/api/probes/batch-ssl-status', methods=['POST'])
def batch_probe_ssl_status():
    data = request.json or {}
    probe_ids = data.get("probe_ids", [])
    
    servers_list = read_servers()
    target_probes = [s for s in servers_list if s["id"] in probe_ids and s.get("os_type") != "windows"]
    
    statuses = {}
    
    def check_single_probe(p):
        pid = p["id"]
        host = p.get("host")
        port = int(p.get("port", 22))
        username = p.get("username")
        password = decrypt_password(p.get("password", ""))
        key_path = p.get("key_path")
        
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        try:
            if key_path and os.path.exists(key_path):
                key = None
                try: key = paramiko.RSAKey.from_private_key_file(key_path)
                except:
                    try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                    except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                ssh.connect(host, port=port, username=username, pkey=key, timeout=2.5, banner_timeout=2.5, auth_timeout=2.5, disabled_algorithms=None)
            else:
                ssh.connect(host, port=port, username=username, password=password, timeout=2.5, banner_timeout=2.5, auth_timeout=2.5, disabled_algorithms=None)
            
            cmd = "grep -m 1 -i '\"EnableSSL\"' /PIXIS/Probe_Daemon/Config/ProbeDaemonSetting.cfg 2>/dev/null || true"
            stdin, stdout, stderr = ssh.exec_command(cmd, timeout=3)
            out = stdout.read().decode('utf-8', errors='ignore').strip()
            ssh.close()
            
            m = re.search(r'"EnableSSL"\s*:\s*(true|false)', out, re.IGNORECASE)
            if m:
                return pid, m.group(1).lower()
            return pid, "未知"
        except Exception:
            return pid, "離線/逾時"

    with concurrent.futures.ThreadPoolExecutor(max_workers=min(10, len(target_probes) or 1)) as executor:
        futures = [executor.submit(check_single_probe, p) for p in target_probes]
        for f in concurrent.futures.as_completed(futures):
            pid, st = f.result()
            statuses[pid] = st

    return jsonify({"success": True, "statuses": statuses})

@app.route('/api/servers/<server_id>/probe-ssl-files', methods=['GET'])
def get_probe_ssl_files(server_id):
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server or server.get("os_type") == "windows":
        return jsonify({"success": False, "message": "找不到該 Probe (Linux) 設備"}), 404

    host = server.get("host")
    port = int(server.get("port", 22))
    username = server.get("username")
    password = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")

    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(host, port=port, username=username, pkey=key, timeout=10, banner_timeout=10, auth_timeout=10, disabled_algorithms=None)
        else:
            ssh.connect(host, port=port, username=username, password=password, timeout=10, banner_timeout=10, auth_timeout=10, disabled_algorithms=None)

        sftp = ssh.open_sftp()
        found_files = []

        # 1. 檢查 Daemon
        daemon_cfg = "/PIXIS/Probe_Daemon/Config/ProbeDaemonSetting.cfg"
        try:
            sftp.stat(daemon_cfg)
            found_files.append(daemon_cfg)
        except Exception:
            pass

        # 2. 自動列出所有 Probe_eth*
        try:
            pixis_dirs = sftp.listdir("/PIXIS")
            eth_dirs = [d for d in pixis_dirs if re.match(r'^Probe_eth\d+$', d)]
            eth_dirs.sort(key=lambda x: int(re.search(r'\d+', x).group()))

            for eth in eth_dirs:
                worker_cfg = f"/PIXIS/{eth}/Config/PortWorkerSetting.cfg"
                try:
                    sftp.stat(worker_cfg)
                    found_files.append(worker_cfg)
                except Exception:
                    pass
        except Exception:
            pass

        sftp.close()
        ssh.close()
        return jsonify({"success": True, "files": found_files, "count": len(found_files)})
    except Exception as e:
        return jsonify({"success": False, "message": f"掃描設定檔失敗: {str(e)}"}), 500

@app.route('/api/servers/<server_id>/probe-ssl', methods=['POST'])
def update_probe_ssl(server_id):
    data = request.json or {}
    enable_ssl = data.get("enable_ssl")
    if enable_ssl is None:
        return jsonify({"success": False, "message": "缺少 enable_ssl 參數"}), 400
    
    target_bool_str = "true" if enable_ssl is True or str(enable_ssl).lower() == "true" else "false"

    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server or server.get("os_type") == "windows":
        return jsonify({"success": False, "message": "找不到該 Probe (Linux) 設備"}), 404

    host = server.get("host")
    port = int(server.get("port", 22))
    username = server.get("username")
    password = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")

    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(host, port=port, username=username, pkey=key, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)
        else:
            ssh.connect(host, port=port, username=username, password=password, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)

        sftp = ssh.open_sftp()
        modified_files = []
        errors = []

        # 1. 處理 Probe_Daemon 設定檔
        daemon_cfg = "/PIXIS/Probe_Daemon/Config/ProbeDaemonSetting.cfg"
        try:
            with sftp.open(daemon_cfg, 'r') as f:
                content = f.read().decode('utf-8', errors='ignore')
            new_content, count = re.subn(r'("EnableSSL"\s*:\s*)(true|false)', rf'\g<1>{target_bool_str}', content, flags=re.IGNORECASE)
            if count > 0:
                with sftp.open(daemon_cfg, 'w') as f:
                    f.write(new_content)
                modified_files.append(daemon_cfg)
            else:
                errors.append(f"{daemon_cfg} 內未找到 EnableSSL 欄位")
        except Exception as e:
            errors.append(f"{daemon_cfg} 讀寫失敗: {str(e)}")

        # 2. 自動偵測 /PIXIS 底下所有 Probe_eth* 並修改 PortWorkerSetting.cfg
        try:
            pixis_dirs = sftp.listdir("/PIXIS")
            eth_dirs = [d for d in pixis_dirs if re.match(r'^Probe_eth\d+$', d)]
            eth_dirs.sort(key=lambda x: int(re.search(r'\d+', x).group()))

            for eth in eth_dirs:
                worker_cfg = f"/PIXIS/{eth}/Config/PortWorkerSetting.cfg"
                try:
                    with sftp.open(worker_cfg, 'r') as f:
                        content = f.read().decode('utf-8', errors='ignore')
                    new_content, count = re.subn(r'("EnableSSL"\s*:\s*)(true|false)', rf'\g<1>{target_bool_str}', content, flags=re.IGNORECASE)
                    if count > 0:
                        with sftp.open(worker_cfg, 'w') as f:
                            f.write(new_content)
                        modified_files.append(worker_cfg)
                    else:
                        errors.append(f"{worker_cfg} 內未找到 EnableSSL 欄位")
                except Exception as e:
                    errors.append(f"{worker_cfg} 讀寫失敗: {str(e)}")
        except Exception as e:
            errors.append(f"列出 /PIXIS 目錄失敗: {str(e)}")

        sftp.close()
        ssh.close()

        if modified_files:
            file_names = ", ".join([os.path.basename(os.path.dirname(os.path.dirname(p))) + "/" + os.path.basename(p) for p in modified_files])
            msg = f"SSL 通訊設定已成功更新為 {target_bool_str} (共更新 {len(modified_files)} 個設定檔: {file_names})"
            return jsonify({"success": True, "message": msg, "modified": modified_files, "errors": errors})
        else:
            return jsonify({"success": False, "message": "未能更新任何設定檔: " + "; ".join(errors)}), 500

    except Exception as e:
        return jsonify({"success": False, "message": f"連線或修改失敗: {str(e)}"}), 500

@app.route('/api/configure-nic-local', methods=['POST'])
def configure_nic_local():
    import base64
    data = request.json
    nic_name = data.get("nic_name")
    ip_address = data.get("ip_address")
    prefix_length = int(data.get("prefix_length", 8))
    
    if not nic_name or not ip_address:
        return jsonify({"success": False, "message": "請輸入網卡名稱與 IP 位址"}), 400
        
    # 定義臨時結果檔案路徑 (寫在 Windows Temp 目錄下)
    temp_dir = os.environ.get("TEMP", "C:\\Windows\\Temp")
    temp_res_file = os.path.join(temp_dir, "nic_res.txt")
    
    # 清理舊 of 臨時檔
    if os.path.exists(temp_res_file):
        try: os.remove(temp_res_file)
        except: pass
            
    # 轉為 PowerShell 相容的正斜線路徑，避免反斜線 (\U, \t 等) 轉義問題
    temp_res_file_ps = temp_res_file.replace("\\", "/")
    
    # 這是要直接在背景執行的 PowerShell 指令
    ps_code = f"""$nicName = "{nic_name}"
$nic = Get-NetAdapter | Where-Object {{ $_.Name -eq $nicName }}
if (-not $nic) {{
    "ERR:找不到名稱為 {nic_name} 的網卡" | Out-File -FilePath "{temp_res_file_ps}" -Encoding utf8
}} else {{
    Set-NetIPInterface -InterfaceIndex $nic.ifIndex -Dhcp Disabled -ErrorAction SilentlyContinue
    Remove-NetIPAddress -InterfaceIndex $nic.ifIndex -Confirm:$false -ErrorAction SilentlyContinue
    $res = New-NetIPAddress -InterfaceIndex $nic.ifIndex -IPAddress "{ip_address}" -PrefixLength {prefix_length} -ErrorAction SilentlyContinue
    if ($res) {{
        "SUCCESS:本地設定完成！" | Out-File -FilePath "{temp_res_file_ps}" -Encoding utf8
    }} else {{
        "ERR:修改 IP 失敗，請確認該網卡是否支援該 IP" | Out-File -FilePath "{temp_res_file_ps}" -Encoding utf8
    }}
}}
"""

    try:
        # 將 PowerShell 程式碼轉換為 Base64 (UTF-16LE 編碼)
        utf16_bytes = ps_code.encode('utf-16-le')
        b64_cmd = base64.b64encode(utf16_bytes).decode('utf-8')
        
        # 檢查目前伺服器程序是否已經具有系統管理員權限
        import ctypes
        is_admin = False
        try:
            is_admin = ctypes.windll.shell32.IsUserAnAdmin() != 0
        except:
            pass
            
        run_res = None
        if is_admin:
            # 如果已經是系統管理員，直接以當前權限在背景執行，完全不會跳出 UAC 確認視窗
            run_res = subprocess.run(
                ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-EncodedCommand", b64_cmd],
                capture_output=True,
                text=True,
                encoding='utf-8',
                errors='ignore',
                shell=True
            )
        else:
            # 若不是管理員，則呼叫 Start-Process 提權執行 EncodedCommand (會彈出 UAC 確認視窗)
            ps_cmd = f"Start-Process powershell -ArgumentList '-NoProfile -ExecutionPolicy Bypass -EncodedCommand {b64_cmd}' -Verb RunAs -WindowStyle Hidden"
            
            run_res = subprocess.run(
                ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps_cmd],
                capture_output=True,
                text=True,
                encoding='utf-8',
                errors='ignore',
                shell=True
            )
        
        # 輪詢等待寫入結果 (最多等 8 秒)
        output = ""
        for _ in range(40):
            if os.path.exists(temp_res_file):
                time.sleep(0.2)  # 確保檔案寫入完整性
                try:
                    with open(temp_res_file, "r", encoding="utf-8-sig") as f:
                        output = f.read().strip()
                    break
                except:
                    pass
            time.sleep(0.2)
            
        # 讀取完畢後清理臨時檔
        if os.path.exists(temp_res_file):
            try: os.remove(temp_res_file)
            except: pass
            
        if not output:
            return jsonify({"success": False, "message": "未偵測到修改結果。若您在 UAC 提權提示中選擇了「否」，則無法完成修改。"}), 400
            
        if output.startswith("SUCCESS:"):
            success_msg = output.split("SUCCESS:")[1].strip()
            return jsonify({"success": True, "message": success_msg})
        elif output.startswith("ERR:"):
            err_msg = output.split("ERR:")[1].strip()
            return jsonify({"success": False, "message": err_msg})
        else:
            return jsonify({"success": False, "message": f"未預期結果: {output}"}), 500
            
    except Exception as e:
        return jsonify({"success": False, "message": f"本地提權執行出錯: {str(e)}"}), 500

def calculate_probe_id(mac) -> int:
    import hashlib
    import struct
    try:
        h = 0
        mac = str.upper(mac).replace(' ','').replace(':','').replace('-','').replace('.','')
        buffer = hashlib.md5(mac.encode('utf-16le')).digest()
        for i in range(4):
            h += struct.unpack('<I', buffer[i * 4 : (i + 1) * 4])[0]
        return h
    except Exception as e:
        print(f"calculate_probe_id error: {e}")
        return None

@app.route('/api/test-probe-port', methods=['POST'])
def test_probe_port():
    import socket
    data = request.json
    probe_index = int(data.get("probe_index", 1))
    port = int(data.get("port", 80))
    local_ip = data.get("local_ip")
    target_ip = data.get("target_ip")
    
    if probe_index < 1 or probe_index > 6:
        return jsonify({"success": False, "message": "無效的 Port 編號"}), 400
        
    ip = target_ip if target_ip else f"{probe_index}.{probe_index}.{probe_index}.{probe_index}"
    username = data.get("username", "root")
    password = data.get("password", "1111")
    
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    sock = None
    try:
        if local_ip:
            # 建立自訂 Socket 并嘗試綁定本機指定的網卡 IP
            sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            sock.settimeout(5)
            
            bind_success = False
            last_bind_err = None
            
            # 輪詢重試機制：最多嘗試 10 次，每次間隔 0.5 秒，總共最多等待 5 秒
            for attempt in range(10):
                try:
                    sock.bind((local_ip, 0))
                    bind_success = True
                    break  # 綁定成功，立刻跳出迴圈
                except Exception as bind_err:
                    last_bind_err = bind_err
                    time.sleep(0.5)  # 失敗則等待 0.5 秒後重試
            
            if not bind_success:
                return jsonify({
                    "success": False,
                    "message": f"無法綁定至本機 IP {local_ip}，請確認該 IP 是否已正確配置在您的網卡上：\n[{type(last_bind_err).__name__}] {str(last_bind_err)}"
                })
                
            sock.connect((ip, 22))
            ssh.connect(ip, port=22, username=username, password=password, sock=sock, timeout=5, disabled_algorithms=None)
        else:
            ssh.connect(ip, port=22, username=username, password=password, timeout=5, disabled_algorithms=None)
    except Exception as e:
        return jsonify({
            "success": False, 
            "message": f"無法 SSH 連線至 Port {probe_index} ({ip}) [綁定本機IP: {local_ip or '無'}]：\n{str(e)}"
        })
        
    try:
        # 取得 MAC 並計算 ProbeID
        mac_cmd = "cat /sys/class/net/eth0/address 2>/dev/null || ip link show eth0 2>/dev/null | awk '/ether/ {print $2}' || ifconfig eth0 2>/dev/null | awk '/ether/ {print $2}'"
        mac_stdin, mac_stdout, mac_stderr = ssh.exec_command(mac_cmd)
        mac_address = mac_stdout.read().decode('utf-8', errors='ignore').strip()
        
        probe_id = None
        if mac_address:
            probe_id = calculate_probe_id(mac_address)

        # 在遠端 Port 上測試指定的 Port (127.0.0.1 本機) 是否開啟與監聽中
        cmd = f"python3 -c \"import socket; s=socket.socket(); s.settimeout(1); s.connect(('127.0.0.1', {port}))\" 2>/dev/null && echo \"OK\" || (python -c \"import socket; s=socket.socket(); s.settimeout(1); s.connect(('127.0.0.1', {port}))\" 2>/dev/null && echo \"OK\" || echo \"FAIL\")"
        stdin, stdout, stderr = ssh.exec_command(cmd)
        result = stdout.read().decode('utf-8', errors='ignore').strip()
        
        if result == "OK":
            return jsonify({
                "success": True,
                "probe_id": probe_id,
                "message": f"測試正常！\nPort {probe_index} ({ip}) 上的 Port {port} 正在監聽中 (OK)"
            })
        else:
            return jsonify({
                "success": False,
                "probe_id": probe_id,
                "message": f"測試失敗！\nPort {probe_index} ({ip}) 上的 Port {port} 未開啟或沒有在監聽 (FAIL)"
            })
    except Exception as e:
        return jsonify({
            "success": False,
            "message": f"在遠端 Port 上執行測試指令時出錯：\n{str(e)}"
        })
    finally:
        ssh.close()

# WebSSH WebSocket 路由
@sock.route('/ws/ssh/<server_id>')
def ws_ssh(ws, server_id):
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        try:
            ws.close(1008, "Server not found")
        except Exception:
            pass
        return
        
    ip = server.get("host")
    port = int(server.get("port", 22))
    username = server.get("username")
    password = server.get("password", "")
    key_path = server.get("key_path", "")
    
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    chan = None
    try:
        if key_path and os.path.exists(key_path):
            key = paramiko.RSAKey.from_private_key_file(key_path)
            ssh.connect(ip, port=port, username=username, pkey=key, timeout=10, disabled_algorithms=None)
        else:
            ssh.connect(ip, port=port, username=username, password=password, timeout=10, disabled_algorithms=None)
            
        # 建立 PTY 交互式 Shell 頻道，寬高設為 xterm.js 常規尺寸
        chan = ssh.invoke_shell(term='xterm', width=100, height=28)
        
        # 建立背景 thread 用於將 SSH 的輸出寫入 WebSocket
        def forward_ssh_to_ws():
            while not chan.closed:
                try:
                    if chan.recv_ready():
                        data = chan.recv(8192)
                        if not data:
                            break
                        ws.send(data.decode('utf-8', errors='ignore'))
                    time.sleep(0.02)
                except Exception:
                    break
            try:
                ws.close()
            except Exception:
                pass
                
        t = threading.Thread(target=forward_ssh_to_ws, daemon=True)
        t.start()
        
        # 主會話循環，將從 WebSocket 接收的輸入寫入 SSH 通道
        while True:
            message = ws.receive()
            if message is None:
                break
            
            is_resize = False
            if isinstance(message, str):
                stripped = message.strip()
                if stripped.startswith('{') and stripped.endswith('}'):
                    try:
                        import json
                        resize_data = json.loads(message)
                        if isinstance(resize_data, dict) and resize_data.get("type") == "resize":
                            cols = int(resize_data.get("cols", 80))
                            rows = int(resize_data.get("rows", 24))
                            chan.resize_pty(width=cols, height=rows)
                            is_resize = True
                    except Exception:
                        pass
            
            if not is_resize:
                chan.send(message)
            
    except Exception as e:
        try:
            ws.send(f"\r\n[SSH 連線錯誤] {str(e)}\r\n")
        except Exception:
            pass
    finally:
        if chan:
            try:
                chan.close()
            except Exception:
                pass
        try:
            ssh.close()
        except Exception:
            pass

def format_file_size(size_bytes):
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    else:
        return f"{size_bytes / (1024 * 1024):.1f} MB"

# ==========================================
# 探針目錄與檔案瀏覽及安全檢查
# ==========================================
def is_safe_probe_path(path):
    if not path:
        return False
    # 防止目錄遍歷與 Windows 反斜線繞過
    if ".." in path or "\\" in path:
        return False
    # 只要是正斜線開頭的 Linux 絕對路徑即可瀏覽與存取
    normalized = os.path.normpath(path).replace('\\', '/')
    if normalized.startswith("/"):
        return True
    return False

@app.route('/api/servers/<server_id>/probe-files', methods=['GET'])
def list_probe_files(server_id):
    path = request.args.get("path", "root").strip()
    
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
        
    if server.get("os_type", "debian") != "debian":
        return jsonify({"success": False, "message": "該伺服器不是 Probe 設備，不支援 SFTP 目錄瀏覽"}), 400

    allowed_bases = [
        "/PIXIS/Probe_Daemon/Logs",
        "/PIXIS/Probe_eth0/Logs",
        "/PIXIS/Probe_eth1/Logs",
        "/PIXIS/Probe_eth2/Logs",
        "/PIXIS/Probe_eth3/Logs"
    ]
    
    # 預設連線初始目錄為 /PIXIS，若使用者點選指定目錄 (如 /root 或 /) 則精確開啟
    if path == "root" or not path:
        path = "/PIXIS"
        
    # 安全檢查
    if not is_safe_probe_path(path):
        return jsonify({"success": False, "message": "不合法的瀏覽路徑，拒絕訪問"}), 403

    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    port = int(server.get("port", 22))
    pwd = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")
    
    files_list = []
    sftp = None
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(server.get("host"), port=port, username=server.get("username"), pkey=key, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)
        else:
            ssh.connect(server.get("host"), port=port, username=server.get("username"), password=pwd, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)
            
        sftp = ssh.open_sftp()
        
        try:
            for attr in sftp.listdir_attr(path):
                filename = attr.filename
                if filename in [".", ".."]:
                    continue
                import stat
                is_dir = stat.S_ISDIR(attr.st_mode)
                permissions_str = stat.filemode(attr.st_mode)
                
                # 解析 owner/group
                owner = "root"
                if attr.longname:
                    parts = attr.longname.split()
                    if len(parts) >= 4:
                        owner = parts[2]
                
                mtime_raw = attr.st_mtime
                mtime_str = datetime.fromtimestamp(mtime_raw).strftime("%Y-%m-%d %H:%M:%S")
                size_str = format_file_size(attr.st_size) if not is_dir else "-"
                
                files_list.append({
                    "name": filename,
                    "type": "directory" if is_dir else "file",
                    "path": f"{path}/{filename}".replace("//", "/"),
                    "mtime_raw": mtime_raw,
                    "mtime": mtime_str,
                    "size": size_str,
                    "size_bytes": attr.st_size if not is_dir else 0,
                    "permissions": permissions_str,
                    "owner": owner
                })
        except FileNotFoundError:
            return jsonify({"success": False, "message": "找不到該目錄"}), 404
        
        # 排序：資料夾置頂，接著依時間最新降序排序
        def get_sort_key(item):
            type_score = 1 if item["type"] == "directory" else 0
            mtime_score = item.get("mtime_raw", 0)
            return (type_score, mtime_score)
            
        files_list.sort(key=get_sort_key, reverse=True)
        return jsonify({"success": True, "current_path": path, "files": files_list})
        
    except Exception as e:
        return jsonify({"success": False, "message": f"SFTP 讀取失敗: {str(e)}"}), 500
    finally:
        if sftp:
            try: sftp.close()
            except: pass
        try: ssh.close()
        except: pass

@app.route('/api/servers/<server_id>/download-probe-log', methods=['GET'])
def download_probe_log(server_id):
    remote_path = request.args.get("path")
    if not remote_path:
        return "缺少路徑參數", 400
        
    if not is_safe_probe_path(remote_path):
        return "不合法的日誌路徑，拒絕讀取", 403
        
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return "找不到該伺服器", 404
        
    filename = os.path.basename(remote_path)
    
    folder = "Probe"
    for part in ["Probe_Daemon", "Probe_eth0", "Probe_eth1", "Probe_eth2", "Probe_eth3"]:
        if part in remote_path:
            folder = part
            break
            
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    port = int(server.get("port", 22))
    pwd = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")
    
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(server.get("host"), port=port, username=server.get("username"), pkey=key, timeout=10, disabled_algorithms=None)
        else:
            ssh.connect(server.get("host"), port=port, username=server.get("username"), password=pwd, timeout=10, disabled_algorithms=None)
            
        sftp = ssh.open_sftp()
        
        try:
            stat_info = sftp.stat(remote_path)
            import stat
            if stat.S_ISDIR(stat_info.st_mode):
                sftp.close()
                ssh.close()
                return "此路徑為目錄，無法直接下載", 400
        except FileNotFoundError:
            sftp.close()
            ssh.close()
            return "伺服器上找不到此檔案", 404
            
        file_stream = io.BytesIO()
        try:
            remote_file = sftp.open(remote_path, "rb")
            file_stream.write(remote_file.read())
            remote_file.close()
        except Exception as file_err:
            sftp.close()
            ssh.close()
            return f"讀取檔案錯誤: {str(file_err)}", 500
            
        sftp.close()
        ssh.close()
        
        file_stream.seek(0)
        
        ip_safe = server.get("host").replace(".", "_")
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        download_name = f"Probe_{ip_safe}_{folder}_{filename.replace('.txt', '').replace('.log', '')}_{timestamp}{os.path.splitext(filename)[1]}"
        
        return send_file(
            file_stream,
            as_attachment=True,
            download_name=download_name,
            mimetype="application/octet-stream"
        )
    except Exception as e:
        return f"下載失敗: {str(e)}", 500

# 暫存 Zip 檔案快取
TEMP_ZIP_CACHE = {}

@app.route('/api/servers/batch-download-probe-logs-stream', methods=['POST'])
def batch_download_probe_logs_stream():
    data = request.json or {}
    days = int(data.get("days", 5))
    server_ids = data.get("server_ids", [])
    basket_items = data.get("basket_items", [])
    
    all_servers = read_servers()
    probe_servers = [s for s in all_servers if s.get("os_type") == "debian" and s.get("category") != "__probe_test__"]
    
    if server_ids and len(server_ids) > 0:
        target_servers = [s for s in probe_servers if s["id"] in server_ids]
    else:
        if basket_items:
            basket_sids = {item.get("server_id") for item in basket_items if item.get("server_id")}
            target_servers = [s for s in probe_servers if s["id"] in basket_sids]
        else:
            target_servers = probe_servers

    def generate_events():
        if not target_servers:
            err_msg = json.dumps({"type": "error", "message": "沒有可進行 Log 下載的 Probe 設備"})
            yield f"data: {err_msg}\n\n"
            return

        now = datetime.now()
        target_dates = set()
        for i in range(days + 1):
            dt = now - timedelta(days=i)
            target_dates.add(dt.strftime("%Y%m%d"))
            
        allowed_base_dirs = [
            "/PIXIS/Probe_Daemon/Logs",
            "/PIXIS/Probe_eth0/Logs",
            "/PIXIS/Probe_eth1/Logs",
            "/PIXIS/Probe_eth2/Logs",
            "/PIXIS/Probe_eth3/Logs",
            "/PIXIS/Logs"
        ]
        
        import tempfile
        import uuid
        zip_id = str(uuid.uuid4())
        temp_dir = os.path.join(tempfile.gettempdir(), "probe_logs")
        os.makedirs(temp_dir, exist_ok=True)
        zip_path = os.path.join(temp_dir, f"{zip_id}.zip")
        
        zf = zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED)
        TEMP_ZIP_CACHE[zip_id] = zip_path
        
        highlights = []
        failed_servers = []
        total_files = 0
        
        for server in target_servers:
            server_id = server["id"]
            server_name = (server.get("name") or "").strip()
            server_host = (server.get("host") or "127.0.0.1").strip()
            if not server_name or server_name == server_host:
                server_folder_name = server_host
            elif server_host in server_name:
                server_folder_name = server_name
            else:
                server_folder_name = f"{server_host}_({server_name})"
                
            msg = json.dumps({"type": "log", "text": f"[BATCH LOG] [{server_host}] 正在建立 SSH/SFTP 連線..."})
            yield f"data: {msg}\n\n"
            
            manual_paths = [item["path"] for item in basket_items if item.get("server_id") == server_id and item.get("path")]
            port = int(server.get("port", 22))
            pwd = decrypt_password(server.get("password", ""))
            key_path = server.get("key_path")
            
            ssh = paramiko.SSHClient()
            ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
            
            try:
                if key_path and os.path.exists(key_path):
                    key = None
                    try: key = paramiko.RSAKey.from_private_key_file(key_path)
                    except:
                        try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                        except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                    ssh.connect(server_host, port=port, username=server.get("username"), pkey=key, timeout=10, disabled_algorithms=None)
                else:
                    ssh.connect(server_host, port=port, username=server.get("username"), password=pwd, timeout=10, disabled_algorithms=None)
                    
                sftp = ssh.open_sftp()
                
                server_files = []
                if manual_paths:
                    for m_path in manual_paths:
                        try:
                            rel_p = m_path
                            if rel_p.startswith("/PIXIS/"): rel_p = rel_p[len("/PIXIS/"):]
                            elif rel_p.startswith("/"): rel_p = rel_p[1:]
                            arc_name = f"{server_folder_name}/{rel_p}"
                            server_files.append((m_path, arc_name))
                        except: pass
                else:
                    msg = json.dumps({"type": "log", "text": f"[BATCH LOG] [{server_host}] 正在掃描 Log 目錄 (近 {days} 天)..."})
                    yield f"data: {msg}\n\n"
                    
                    def scan_dir(dir_path, depth=0):
                        if depth > 1: return []
                        found_list = []
                        try:
                            for attr in sftp.listdir_attr(dir_path):
                                fn = attr.filename
                                if fn in [".", ".."]: continue
                                full_p = f"{dir_path}/{fn}".replace("//", "/")
                                import stat
                                if stat.S_ISDIR(attr.st_mode):
                                    found_list.extend(scan_dir(full_p, depth + 1))
                                else:
                                    found_list.append((full_p, fn, attr.st_mtime, attr.st_size))
                        except: pass
                        return found_list
                        
                    for b_dir in allowed_base_dirs:
                        all_files_in_bdir = scan_dir(b_dir)
                        if not all_files_in_bdir: continue
                        selected_paths_in_bdir = set()
                        for f_path, fn, mtime, sz in all_files_in_bdir:
                            fn_lower = fn.lower()
                            if fn_lower.endswith('.pcap') or fn_lower.endswith('.pcapng') or fn_lower.endswith('.tar') or fn_lower.endswith('.zip') or fn_lower.endswith('.gz'):
                                continue
                            is_valid_log = (fn_lower.endswith('.log') or fn_lower.endswith('.txt') or 'eventlogs' in fn_lower or 'daemonevent' in fn_lower or 'dhcpservice' in fn_lower)
                            if not is_valid_log: continue
                            if sz > 50 * 1024 * 1024:
                                highlights.append({"server_name": server_name, "host": server_host, "file": fn, "reason": f"檔案過大 ({format_file_size(sz)} > 50MB)，已自動跳過"})
                                continue
                            m = re.search(r'(\d{8})', fn)
                            if m:
                                date_str = m.group(1)
                                if date_str in target_dates:
                                    selected_paths_in_bdir.add(f_path)
                            else:
                                m_date_str = datetime.fromtimestamp(mtime).strftime("%Y%m%d")
                                if m_date_str in target_dates:
                                    selected_paths_in_bdir.add(f_path)
                        
                        for f_path, fn, mtime, sz in all_files_in_bdir:
                            if f_path in selected_paths_in_bdir:
                                rel_p = f_path
                                if rel_p.startswith("/PIXIS/"): rel_p = rel_p[len("/PIXIS/"):]
                                elif rel_p.startswith("/"): rel_p = rel_p[1:]
                                arc_name = f"{server_folder_name}/{rel_p}"
                                server_files.append((f_path, arc_name))
                                
                for r_path, arc_n in server_files:
                    try:
                        st_size = 0
                        try:
                            st_info = sftp.stat(r_path)
                            st_size = st_info.st_size
                        except: pass
                        
                        rf = sftp.open(r_path, "rb")
                        if st_size > 0: content = rf.read(st_size)
                        else: content = rf.read()
                        rf.close()
                        
                        zf.writestr(arc_n, content)
                        total_files += 1
                        
                        msg_txt = f"[BATCH LOG] [{server_host}] [成功] 讀取檔案: {r_path} ({format_file_size(st_size)})"
                        yield f"data: {json.dumps({'type': 'log', 'text': msg_txt})}\n\n"
                    except Exception as file_ex:
                        err_txt = f"[BATCH LOG] [{server_host}] [失敗] 讀取檔案失敗: {r_path} ({file_ex})"
                        yield f"data: {json.dumps({'type': 'log', 'text': err_txt})}\n\n"
                        
                sftp.close()
                ssh.close()
            except Exception as e:
                err_conn = f"[BATCH LOG] [{server_host}] [連線失敗/超時]: {e}"
                yield f"data: {json.dumps({'type': 'log', 'text': err_conn})}\n\n"
                failed_servers.append({"host": server_host, "name": server_name, "reason": f"連線失敗/超時 ({str(e)})"})
                highlights.append({"server_name": server_name, "host": server_host, "file": "連線失敗", "reason": f"連線超時或無法連接 ({str(e)})"})
                try: ssh.close()
                except: pass

        zf.close()
        
        comp_msg = json.dumps({
            "type": "complete",
            "download_url": f"/api/servers/download-temp-zip/{zip_id}",
            "total_files": total_files,
            "highlights": highlights,
            "failed_servers": failed_servers
        })
        yield f"data: {comp_msg}\n\n"

    return Response(generate_events(), mimetype='text/event-stream')

@app.route('/api/servers/download-temp-zip/<zip_id>', methods=['GET'])
def download_temp_zip(zip_id):
    zip_path = TEMP_ZIP_CACHE.get(zip_id)
    if not zip_path or not os.path.exists(zip_path):
        import tempfile
        temp_dir = os.path.join(tempfile.gettempdir(), "probe_logs")
        zip_path = os.path.join(temp_dir, f"{zip_id}.zip")
        if not os.path.exists(zip_path):
            return "ZIP 檔案不存在或已過期", 404
            
    now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    download_name = f"Batch_Probe_Logs_{now_str}.zip"
    
    return send_file(
        zip_path,
        as_attachment=True,
        download_name=download_name,
        mimetype="application/zip"
    )

@app.route('/api/servers/batch-download-probe-logs', methods=['POST'])
def batch_download_probe_logs():
    data = request.json or {}
    days = int(data.get("days", 5))
    server_ids = data.get("server_ids", [])
    basket_items = data.get("basket_items", []) # [{"server_id": "...", "path": "..."}]
    
    all_servers = read_servers()
    probe_servers = [s for s in all_servers if s.get("os_type") == "debian" and s.get("category") != "__probe_test__"]
    
    if server_ids and len(server_ids) > 0:
        target_servers = [s for s in probe_servers if s["id"] in server_ids]
    else:
        if basket_items:
            basket_sids = {item.get("server_id") for item in basket_items if item.get("server_id")}
            target_servers = [s for s in probe_servers if s["id"] in basket_sids]
        else:
            target_servers = probe_servers
            
    print(f"\n[BATCH LOG] ===== 開始處理 Probe 批次 Log 下載 ===== (天數: {days}, 目標設備數: {len(target_servers)}, 手動勾選數: {len(basket_items)})")
    
    if not target_servers:
        print("[BATCH LOG] [錯誤] 沒有找到任何可進行 Log 下載的 Probe 設備！")
        return jsonify({"success": False, "message": "沒有可進行 Log 下載的 Probe 設備"}), 400
        
    now = datetime.now()
    target_dates = set()
    for i in range(days + 1):
        dt = now - timedelta(days=i)
        target_dates.add(dt.strftime("%Y%m%d"))
        
    allowed_base_dirs = [
        "/PIXIS/Probe_Daemon/Logs",
        "/PIXIS/Probe_eth0/Logs",
        "/PIXIS/Probe_eth1/Logs",
        "/PIXIS/Probe_eth2/Logs",
        "/PIXIS/Probe_eth3/Logs",
        "/PIXIS/Logs"
    ]
    
    zip_buffer = io.BytesIO()
    highlights = []
    
    def process_server_logs(server):
        server_files = []
        server_id = server["id"]
        server_name = (server.get("name") or "").strip()
        server_host = (server.get("host") or "127.0.0.1").strip()
        if not server_name or server_name == server_host:
            server_folder_name = server_host
        elif server_host in server_name:
            server_folder_name = server_name
        else:
            server_folder_name = f"{server_host}_({server_name})"
        
        manual_paths = [item["path"] for item in basket_items if item.get("server_id") == server_id and item.get("path")]
                
        port = int(server.get("port", 22))
        pwd = decrypt_password(server.get("password", ""))
        key_path = server.get("key_path")
        
        print(f"[BATCH LOG] [{server_host} ({server_name})] 正在進行 SSH 連線 (Port: {port})...")
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        
        try:
            if key_path and os.path.exists(key_path):
                key = None
                try: key = paramiko.RSAKey.from_private_key_file(key_path)
                except:
                    try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                    except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                ssh.connect(server_host, port=port, username=server.get("username"), pkey=key, timeout=10, disabled_algorithms=None)
            else:
                ssh.connect(server_host, port=port, username=server.get("username"), password=pwd, timeout=10, disabled_algorithms=None)
                
            sftp = ssh.open_sftp()
            print(f"[BATCH LOG] [{server_host}] SSH/SFTP 連線成功！")
            
            # 【Fast-Path 1】：若使用者手動勾選了特定的檔案，直接下載這幾個檔案
            if manual_paths:
                print(f"[BATCH LOG] [{server_host}] 觸發手動勾選 Fast-Path (共 {len(manual_paths)} 個指定檔案)")
                for m_path in manual_paths:
                    try:
                        rel_p = m_path
                        if rel_p.startswith("/PIXIS/"):
                            rel_p = rel_p[len("/PIXIS/"):]
                        elif rel_p.startswith("/"):
                            rel_p = rel_p[1:]
                        arc_name = f"{server_folder_name}/{rel_p}"
                        server_files.append((m_path, arc_name))
                    except Exception as m_err:
                        print(f"[BATCH LOG] [{server_host}] 手動路徑處理失敗 ({m_path}): {m_err}")
            else:
                # 【Fast-Path 2】：若無手動勾選特定檔案，才執行一鍵 N 天的自動目錄掃描
                print(f"[BATCH LOG] [{server_host}] 觸發天數自動掃描 (目標天數: {days} 天)")
                def scan_dir(dir_path, depth=0):
                    if depth > 1:
                        return []
                    found_list = []
                    try:
                        for attr in sftp.listdir_attr(dir_path):
                            fn = attr.filename
                            if fn in [".", ".."]: continue
                            full_p = f"{dir_path}/{fn}".replace("//", "/")
                            import stat
                            if stat.S_ISDIR(attr.st_mode):
                                found_list.extend(scan_dir(full_p, depth + 1))
                            else:
                                found_list.append((full_p, fn, attr.st_mtime, attr.st_size))
                    except Exception as scan_err:
                        pass
                    return found_list
                    
                for b_dir in allowed_base_dirs:
                    print(f"[BATCH LOG] [{server_host}] 掃描目錄: {b_dir}")
                    all_files_in_bdir = scan_dir(b_dir)
                    if not all_files_in_bdir:
                        continue
                        
                    selected_paths_in_bdir = set()
                    
                    for f_path, fn, mtime, sz in all_files_in_bdir:
                        fn_lower = fn.lower()
                        # 排除非日誌檔 (如 .pcap, .pcapng, .tar, .zip, .gz, .iso 等)
                        if fn_lower.endswith('.pcap') or fn_lower.endswith('.pcapng') or fn_lower.endswith('.tar') or fn_lower.endswith('.zip') or fn_lower.endswith('.gz'):
                            print(f"[BATCH LOG] [{server_host}] 跳過非日誌/封包檔: {fn}")
                            continue
                            
                        # 嚴格驗證副檔名與 Probe 日誌類型
                        is_valid_log = (
                            fn_lower.endswith('.log') or 
                            fn_lower.endswith('.txt') or 
                            'eventlogs' in fn_lower or 
                            'daemonevent' in fn_lower or 
                            'dhcpservice' in fn_lower
                        )
                        if not is_valid_log:
                            print(f"[BATCH LOG] [{server_host}] 跳過非 Log/TXT 檔: {fn}")
                            continue
                            
                        # 50MB 大小保護防護
                        if sz > 50 * 1024 * 1024:
                            print(f"[BATCH LOG] [{server_host}] 跳過大於 50MB 檔案: {fn} ({sz} bytes)")
                            highlights.append({
                                "server_name": server_name,
                                "host": server_host,
                                "file": fn,
                                "reason": f"檔案過大 ({format_file_size(sz)} > 50MB)，已自動跳過"
                            })
                            continue

                        m = re.search(r'(\d{8})', fn)
                        if m:
                            date_str = m.group(1)
                            if date_str in target_dates:
                                selected_paths_in_bdir.add(f_path)
                                try:
                                    f_date = datetime.strptime(date_str, "%Y%m%d")
                                    m_date = datetime.fromtimestamp(mtime)
                                    if abs((f_date - m_date).days) > 2:
                                        highlights.append({
                                            "server_name": server_name,
                                            "host": server_host,
                                            "file": fn,
                                            "reason": f"檔名日期({date_str})與實體修改時間({m_date.strftime('%Y-%m-%d')})不符合"
                                        })
                                except: pass
                        else:
                            # 無日期檔名的活檔：檢查修改時間是否在目標天數內
                            m_date_str = datetime.fromtimestamp(mtime).strftime("%Y%m%d")
                            if m_date_str in target_dates:
                                selected_paths_in_bdir.add(f_path)
                            
                    if not selected_paths_in_bdir and all_files_in_bdir:
                        all_files_in_bdir.sort(key=lambda x: x[2], reverse=True)
                        top_files = all_files_in_bdir[:(days + 1)]
                        for f_path, fn, mtime, sz in top_files:
                            selected_paths_in_bdir.add(f_path)
                            m_date_str = datetime.fromtimestamp(mtime).strftime('%Y-%m-%d')
                            highlights.append({
                                "server_name": server_name,
                                "host": server_host,
                                "file": fn,
                                "reason": f"設備時間偏離，已自動保底抓取最新修改日誌 ({m_date_str})"
                            })
                            
                    for f_path, fn, mtime, sz in all_files_in_bdir:
                        if f_path in selected_paths_in_bdir:
                            rel_p = f_path
                            if rel_p.startswith("/PIXIS/"):
                                rel_p = rel_p[len("/PIXIS/"):]
                            elif rel_p.startswith("/"):
                                rel_p = rel_p[1:]
                                
                            arc_name = f"{server_folder_name}/{rel_p}"
                            server_files.append((f_path, arc_name))
                        
            print(f"[BATCH LOG] [{server_host}] 評估完畢，準備讀取並寫入 {len(server_files)} 個目標 Log 檔案...")
            downloaded_bytes = []
            for r_path, arc_n in server_files:
                try:
                    st_size = 0
                    try:
                        st_info = sftp.stat(r_path)
                        st_size = st_info.st_size
                    except: pass
                    
                    f_io = io.BytesIO()
                    rf = sftp.open(r_path, "rb")
                    if st_size > 0:
                        f_io.write(rf.read(st_size))
                    else:
                        f_io.write(rf.read())
                    rf.close()
                    f_io.seek(0)
                    downloaded_bytes.append((arc_n, f_io.getvalue()))
                    print(f"[BATCH LOG] [{server_host}] [成功] 讀取檔案: {r_path} ({st_size} bytes)")
                except Exception as ex:
                    print(f"[BATCH LOG] [{server_host}] [錯誤] 讀取檔案失敗: {r_path}, 原因: {ex}")
                    
            sftp.close()
            ssh.close()
            return downloaded_bytes
        except Exception as e:
            print(f"[BATCH LOG] [{server_host}] [連線失敗/超時]: {e}")
            highlights.append({
                "server_name": server_name,
                "host": server_host,
                "file": "連線失敗",
                "reason": f"連線超時(10秒)或無法連接 ({str(e)})"
            })
            try: ssh.close()
            except: pass
            return []

    results = []
    for s in target_servers:
        results.append(process_server_logs(s))
        
    total_files_added = 0
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for res_list in results:
            for arc_n, data_bytes in res_list:
                zf.writestr(arc_n, data_bytes)
                total_files_added += 1
                
    print(f"[BATCH LOG] ===== 批次打包完成 ===== (總共打包 {total_files_added} 個檔案, 異常記錄數: {len(highlights)})\n")
    
    if total_files_added == 0:
        return jsonify({"success": False, "message": "未找到任何符合條件的 Log 檔案，請檢查設備連線或選擇的天數。"}), 404
        
    zip_buffer.seek(0)
    zip_filename = f"Batch_Probe_Logs_{now.strftime('%Y%m%d_%H%M%S')}.zip"
    
    response = send_file(
        zip_buffer,
        as_attachment=True,
        download_name=zip_filename,
        mimetype="application/zip"
    )
    
    if highlights:
        encoded_hl = urllib.parse.quote(json.dumps(highlights, ensure_ascii=False))
        response.headers["X-Abnormal-Highlights"] = encoded_hl
        response.headers["Access-Control-Expose-Headers"] = "X-Abnormal-Highlights"
        
    return response

def sftp_mkdirs(sftp, remote_directory):
    """遞迴建立 SFTP 目錄"""
    if remote_directory == "/" or remote_directory == "" or remote_directory == ".":
        return
    try:
        sftp.stat(remote_directory)
    except FileNotFoundError:
        parent = os.path.dirname(remote_directory.rstrip('/'))
        sftp_mkdirs(sftp, parent)
        try:
            sftp.mkdir(remote_directory)
        except IOError:
            pass

@app.route('/api/servers/<server_id>/upload-probe-file', methods=['POST'])
def upload_probe_file(server_id):
    path = request.form.get("path")
    relative_path = request.form.get("relative_path")
    if not path:
        return jsonify({"success": False, "message": "缺少目標目錄參數"}), 400
        
    if 'file' not in request.files:
        return jsonify({"success": False, "message": "缺少上傳的檔案"}), 400
        
    uploaded_file = request.files['file']
    filename = uploaded_file.filename
    if not filename:
        return jsonify({"success": False, "message": "無效的檔名"}), 400
        
    # 安全檢查與路徑計算
    if relative_path:
        remote_file_path = f"{path}/{relative_path}".replace("//", "/")
    else:
        remote_file_path = f"{path}/{filename}".replace("//", "/")
        
    if not is_safe_probe_path(remote_file_path):
        return jsonify({"success": False, "message": "不合法或受限的目標路徑，拒絕寫入"}), 403
        
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
        
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    port = int(server.get("port", 22))
    pwd = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")
    
    sftp = None
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(server.get("host"), port=port, username=server.get("username"), pkey=key, timeout=10, disabled_algorithms=None)
        else:
            ssh.connect(server.get("host"), port=port, username=server.get("username"), password=pwd, timeout=10, disabled_algorithms=None)
            
        sftp = ssh.open_sftp()
        
        # 遞迴建立父目錄
        remote_parent_dir = os.path.dirname(remote_file_path)
        sftp_mkdirs(sftp, remote_parent_dir)
        
        # 使用 putfo 直接將記憶體中的檔案資料流傳輸給 SFTP
        file_stream = io.BytesIO(uploaded_file.read())
        sftp.putfo(file_stream, remote_file_path)
        
        return jsonify({"success": True, "message": "檔案上傳替換成功！"})
    except Exception as e:
        return jsonify({"success": False, "message": f"上傳失敗: {str(e)}"}), 500
    finally:
        if sftp:
            try: sftp.close()
            except: pass
        try: ssh.close()
        except: pass

@app.route('/api/servers/batch-upload-probe-file', methods=['POST'])
def batch_upload_probe_file():
    path = request.form.get("path")
    relative_path = request.form.get("relative_path")
    server_ids_json = request.form.get("server_ids")
    if not path or not server_ids_json:
        return jsonify({"success": False, "message": "缺少目標路徑或目標設備清單"}), 400
        
    try:
        server_ids = json.loads(server_ids_json)
    except Exception:
        return jsonify({"success": False, "message": "無效的設備清單格式"}), 400
        
    if 'file' not in request.files:
        return jsonify({"success": False, "message": "缺少上傳的檔案"}), 400
        
    uploaded_file = request.files['file']
    filename = uploaded_file.filename
    if not filename:
        return jsonify({"success": False, "message": "無效的檔名"}), 400
        
    # 安全檢查與路徑計算
    if relative_path:
        remote_file_path = f"{path}/{relative_path}".replace("//", "/")
    else:
        remote_file_path = f"{path}/{filename}".replace("//", "/")
        
    if not is_safe_probe_path(remote_file_path):
        return jsonify({"success": False, "message": "不合法或受限的目標路徑，拒絕寫入"}), 403
        
    # 先把檔案內容讀進記憶體
    file_bytes = uploaded_file.read()
    
    servers = read_servers()
    results = []
    
    # 用於多執行緒執行單個上傳的輔助函數
    def upload_to_single_server(s_id):
        s = next((srv for srv in servers if srv["id"] == s_id), None)
        if not s:
            return {"server_name": "未知設備", "host": "-", "status": "failed", "message": "找不到該設備設定"}
            
        s_name = s.get("name", "未命名")
        s_host = s.get("host", "-")
        
        if s.get("os_type", "debian") != "debian":
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": "該伺服器不是 Debian Probe"}
            
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        port = int(s.get("port", 22))
        pwd = decrypt_password(s.get("password", ""))
        key_path = s.get("key_path")
        
        sftp = None
        try:
            if key_path and os.path.exists(key_path):
                key = None
                try: key = paramiko.RSAKey.from_private_key_file(key_path)
                except:
                    try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                    except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                ssh.connect(s_host, port=port, username=s.get("username"), pkey=key, timeout=8, disabled_algorithms=None)
            else:
                ssh.connect(s_host, port=port, username=s.get("username"), password=pwd, timeout=8, disabled_algorithms=None)
                
            sftp = ssh.open_sftp()
            
            # 遞迴建立父目錄
            remote_parent_dir = os.path.dirname(remote_file_path)
            sftp_mkdirs(sftp, remote_parent_dir)
            
            sftp.putfo(io.BytesIO(file_bytes), remote_file_path)
            return {"server_name": s_name, "host": s_host, "status": "success", "message": "成功"}
        except Exception as e:
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": str(e)}
        finally:
            if sftp:
                try: sftp.close()
                except: pass
            try: ssh.close()
            except: pass
            
    # 使用執行緒池並行上傳，效率極高
    with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(upload_to_single_server, s_id): s_id for s_id in server_ids}
        for future in concurrent.futures.as_completed(futures):
            results.append(future.result())
            
    return jsonify({"success": True, "results": results})

@app.route('/api/servers/<server_id>/reboot', methods=['POST'])
def reboot_server(server_id):
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
        
    if server.get("os_type", "debian") != "debian":
        return jsonify({"success": False, "message": "該伺服器不是 Debian Probe"}), 400
        
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    port = int(server.get("port", 22))
    pwd = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")
    
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(server.get("host"), port=port, username=server.get("username"), pkey=key, timeout=5, disabled_algorithms=None)
        else:
            ssh.connect(server.get("host"), port=port, username=server.get("username"), password=pwd, timeout=5, disabled_algorithms=None)
            
        ssh.exec_command("sudo reboot || reboot")
        time.sleep(0.5)
        ssh.close()
        return jsonify({"success": True, "message": "已成功發送重啟指令！"})
    except Exception as e:
        err_msg = str(e)
        if "SSH session not active" in err_msg or "Connection reset" in err_msg or "EOF" in err_msg or "broken pipe" in err_msg.lower():
            return jsonify({"success": True, "message": "已成功發送重啟指令 (連線已中斷)！"})
        return jsonify({"success": False, "message": f"發送重啟指令失敗: {err_msg}"}), 500

@app.route('/api/servers/batch-reboot', methods=['POST'])
def batch_reboot_servers():
    server_ids = request.json.get("server_ids")
    if not server_ids:
        return jsonify({"success": False, "message": "缺少目標設備清單"}), 400
        
    servers = read_servers()
    results = []
    
    def reboot_single_server(s_id):
        s = next((srv for srv in servers if srv["id"] == s_id), None)
        if not s:
            return {"server_name": "未知設備", "host": "-", "status": "failed", "message": "找不到該設備設定"}
            
        s_name = s.get("name", "未命名")
        s_host = s.get("host", "-")
        
        if s.get("os_type", "debian") != "debian":
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": "該伺服器不是 Debian Probe"}
            
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        port = int(s.get("port", 22))
        pwd = decrypt_password(s.get("password", ""))
        key_path = s.get("key_path")
        
        try:
            if key_path and os.path.exists(key_path):
                key = None
                try: key = paramiko.RSAKey.from_private_key_file(key_path)
                except:
                    try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                    except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                ssh.connect(s_host, port=port, username=s.get("username"), pkey=key, timeout=5, disabled_algorithms=None)
            else:
                ssh.connect(s_host, port=port, username=s.get("username"), password=pwd, timeout=5, disabled_algorithms=None)
                
            ssh.exec_command("sudo reboot || reboot")
            time.sleep(0.5)
            ssh.close()
            return {"server_name": s_name, "host": s_host, "status": "success", "message": "已成功發送重啟指令！"}
        except Exception as e:
            err_msg = str(e)
            if "SSH session not active" in err_msg or "Connection reset" in err_msg or "EOF" in err_msg or "broken pipe" in err_msg.lower():
                return {"server_name": s_name, "host": s_host, "status": "success", "message": "已成功發送重啟指令 (連線已中斷)！"}
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": f"發送重啟指令失敗: {err_msg}"}

    with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(reboot_single_server, s_id): s_id for s_id in server_ids}
        for future in concurrent.futures.as_completed(futures):
            results.append(future.result())
            
    return jsonify({"success": True, "results": results})

@app.route('/api/servers/batch-exec-script', methods=['POST'])
def batch_exec_script():
    script_path = request.json.get("script_path")
    server_ids = request.json.get("server_ids")
    if not script_path or not server_ids:
        return jsonify({"success": False, "message": "缺少腳本路徑或目標設備列表"}), 400

    servers = read_servers()
    results = []

    # 解析腳本目錄與腳本檔名
    script_dir = os.path.dirname(script_path).replace('\\', '/')
    script_name = os.path.basename(script_path)

    def exec_single_server(s_id):
        s = next((srv for srv in servers if srv["id"] == s_id), None)
        if not s:
            return {"server_name": "未知設備", "host": "-", "status": "failed", "message": "找不到該設備設定", "output": ""}

        s_name = s.get("name", "未命名")
        s_host = s.get("host", "-")

        if s.get("os_type", "debian") != "debian":
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": "該伺服器不是 Debian Probe", "output": ""}

        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        port = int(s.get("port", 22))
        pwd = decrypt_password(s.get("password", ""))
        key_path = s.get("key_path")

        try:
            if key_path and os.path.exists(key_path):
                key = None
                try: key = paramiko.RSAKey.from_private_key_file(key_path)
                except:
                    try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                    except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                ssh.connect(s_host, port=port, username=s.get("username"), pkey=key, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)
            else:
                ssh.connect(s_host, port=port, username=s.get("username"), password=pwd, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)

            cmd = f"cd '{script_dir}' && (sed -i 's/\\r$//' '{script_name}' 2>/dev/null || true) && chmod +x '{script_name}' && (printf '\\ny\\nyes\\n\\n' 2>/dev/null | ./'{script_name}')"
            stdin, stdout, stderr = ssh.exec_command(cmd, get_pty=True, timeout=300)
            out_str = stdout.read().decode('utf-8', errors='ignore')
            err_str = stderr.read().decode('utf-8', errors='ignore')
            exit_code = stdout.channel.recv_exit_status()

            full_output = (out_str + ("\n[STDERR]\n" + err_str if err_str else "")).strip()

            if exit_code == 0:
                return {"server_name": s_name, "host": s_host, "status": "success", "message": "執行成功 (exit code: 0)", "output": full_output or "執行完成，無 Console 輸出"}
            else:
                return {"server_name": s_name, "host": s_host, "status": "failed", "message": f"執行失敗 (exit code: {exit_code})", "output": full_output or "未知錯誤"}
        except Exception as e:
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": str(e), "output": ""}
        finally:
            try: ssh.close()
            except: pass

    with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(exec_single_server, s_id): s_id for s_id in server_ids}
        for future in concurrent.futures.as_completed(futures):
            results.append(future.result())

    return jsonify({"success": True, "results": results})

# ---------------------------------------------------------
# Program 腳本資料夾管理與批次部署執行 API
# ---------------------------------------------------------
@app.route('/api/program/files', methods=['GET'])
def get_program_files():
    try:
        program_dir = os.path.join(app.root_path, 'Program')
        os.makedirs(program_dir, exist_ok=True)
        files = []
        for f in os.listdir(program_dir):
            fp = os.path.join(program_dir, f)
            if os.path.isfile(fp):
                st = os.stat(fp)
                files.append({
                    "name": f,
                    "size": st.st_size,
                    "mtime": datetime.fromtimestamp(st.st_mtime).strftime('%Y-%m-%d %H:%M:%S'),
                    "path": fp.replace('\\', '/')
                })
        return jsonify({"success": True, "files": files})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route('/api/servers/batch-program-deploy', methods=['POST'])
def batch_program_deploy():
    filenames = request.json.get("filenames", [])
    server_ids = request.json.get("server_ids", [])
    auto_exec = request.json.get("auto_exec", True)
    target_dir = request.json.get("target_dir", "/root").rstrip('/')

    if not filenames or not server_ids:
        return jsonify({"success": False, "message": "缺少要部署的檔案或目標設備清單"}), 400

    program_dir = os.path.join(app.root_path, 'Program')
    servers = read_servers()
    results = []

    def deploy_to_single_server(s_id):
        s = next((srv for srv in servers if srv["id"] == s_id), None)
        if not s:
            return {"server_name": "未知設備", "host": "-", "status": "failed", "message": "找不到該設備設定", "output": ""}

        s_name = s.get("name", "未命名")
        s_host = s.get("host", "-")

        if s.get("os_type", "debian") != "debian":
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": "該伺服器不是 Debian Probe", "output": ""}

        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        port = int(s.get("port", 22))
        pwd = decrypt_password(s.get("password", ""))
        key_path = s.get("key_path")

        sftp = None
        logs = []
        try:
            if key_path and os.path.exists(key_path):
                key = None
                try: key = paramiko.RSAKey.from_private_key_file(key_path)
                except:
                    try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                    except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
                ssh.connect(s_host, port=port, username=s.get("username"), pkey=key, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)
            else:
                ssh.connect(s_host, port=port, username=s.get("username"), password=pwd, timeout=15, banner_timeout=15, auth_timeout=15, disabled_algorithms=None)

            sftp = ssh.open_sftp()
            sftp_mkdirs(sftp, target_dir)

            uploaded_scripts = []
            for fname in filenames:
                local_path = os.path.join(program_dir, fname)
                if not os.path.exists(local_path):
                    logs.append(f"[警告] 找不到本機 Program 檔案: {fname}")
                    continue
                remote_path = f"{target_dir}/{fname}"
                sftp.put(local_path, remote_path)
                logs.append(f"[成功] 檔案 {fname} 已上傳至 {remote_path}")
                uploaded_scripts.append(fname)

            exec_outputs = []
            if auto_exec and uploaded_scripts:
                for script_name in uploaded_scripts:
                    cmd = f"cd '{target_dir}' && (sed -i 's/\\r$//' '{script_name}' 2>/dev/null || true) && chmod +x '{script_name}' && (printf '\\ny\\nyes\\n\\n' 2>/dev/null | ./'{script_name}')"
                    logs.append(f"[執行] 指令: ./{script_name}")
                    stdin, stdout, stderr = ssh.exec_command(cmd, get_pty=True, timeout=300)
                    out_str = stdout.read().decode('utf-8', errors='ignore')
                    err_str = stderr.read().decode('utf-8', errors='ignore')
                    exit_code = stdout.channel.recv_exit_status()
                    full_out = (out_str + ("\n[STDERR]\n" + err_str if err_str else "")).strip()
                    exec_outputs.append(f"=== {script_name} (exit code: {exit_code}) ===\n{full_out or '無輸出'}")

            final_msg = "\n".join(logs)
            if exec_outputs:
                final_msg += "\n\n--- 腳本執行結果 ---\n" + "\n\n".join(exec_outputs)

            return {"server_name": s_name, "host": s_host, "status": "success", "message": "傳送與執行成功", "output": final_msg}
        except Exception as e:
            return {"server_name": s_name, "host": s_host, "status": "failed", "message": f"部署失敗: {str(e)}", "output": "\n".join(logs)}
        finally:
            if sftp:
                try: sftp.close()
                except: pass
            try: ssh.close()
            except: pass

    with concurrent.futures.ThreadPoolExecutor(max_workers=10) as executor:
        futures = {executor.submit(deploy_to_single_server, s_id): s_id for s_id in server_ids}
        for future in concurrent.futures.as_completed(futures):
            results.append(future.result())

    return jsonify({"success": True, "results": results})


@app.route('/api/local-nics', methods=['GET'])
def get_local_nics():
    try:
        cmd = [
            "powershell",
            "-NoProfile",
            "-Command",
            "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8; Get-NetIPInterface -AddressFamily IPv4 -ConnectionState Connected | Where-Object InterfaceAlias -notlike '*Loopback*' | Get-NetIPAddress | Select-Object InterfaceAlias, IPAddress | ConvertTo-Json"
        ]
        res = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='ignore')
        
        if not res.stdout.strip():
            return jsonify([])
            
        data = json.loads(res.stdout)
        if isinstance(data, dict):
            data = [data]
            
        nics = []
        seen = set()
        for item in data:
            name = item.get("InterfaceAlias")
            ip = item.get("IPAddress")
            if name and ip:
                # 排除 Windows 自動生成的虛擬熱點介面（如「區域連線* 1」）
                if "*" in name:
                    continue
                # 排除其他虛擬與無關網卡（如虛擬機、藍牙等）
                name_lower = name.lower()
                if any(x in name_lower for x in ["loopback", "bluetooth", "藍牙", "vmware", "virtual", "vbox", "tunnel", "wsl", "vethernet"]):
                    continue
                key = (name, ip)
                if key not in seen:
                    seen.add(key)
                    nics.append({"name": name, "ip": ip})
        return jsonify(nics)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/detect-probe-ports', methods=['POST'])
def detect_probe_ports():
    data = request.json or {}
    ip = data.get("ip", "1.1.1.1")
    username = data.get("username", "root")
    password = data.get("password", "1111")
    
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    sftp = None
    try:
        # 連接遠端探針，讀取設定檔
        ssh.connect(ip, port=22, username=username, password=password, timeout=5, disabled_algorithms=None)
        sftp = ssh.open_sftp()
        config_path = "/PIXIS/Probe_Daemon/Config/ProbeDaemonSetting.cfg"
        
        with sftp.open(config_path, "r") as f:
            content = f.read().decode('utf-8', errors='ignore')
            
        # 解析 JSON 獲取 ManagementPortWorkerInfo 的陣列長度
        import json
        config_data = json.loads(content)
        probe_config = config_data.get("ProbeDaemonConfig", {})
        port_workers = probe_config.get("ManagementPortWorkerInfo", [])
        
        port_count = len(port_workers) if isinstance(port_workers, list) else 4
        if port_count == 0:
            port_count = 4
            
        return jsonify({"success": True, "port_count": port_count})
    except Exception as e:
        return jsonify({"success": False, "message": f"無法透過 SFTP 連線至探針並讀取設定檔：{str(e)}"}), 200
    finally:
        if sftp:
            try: sftp.close()
            except: pass
        try: ssh.close()
        except: pass

@app.route('/api/servers/<server_id>/delete-probe-file', methods=['POST'])
def delete_probe_file(server_id):
    path = request.json.get("path")
    if not path:
        return jsonify({"success": False, "message": "缺少目標路徑參數"}), 400
        
    # 安全檢查
    if not is_safe_probe_path(path):
        return jsonify({"success": False, "message": "不合法或受限的路徑，拒絕刪除"}), 403
        
    # 限制：絕對不允許刪除 5 個基礎 Logs 目錄以及 /PIXIS 根目錄本身
    allowed_bases = [
        "/PIXIS",
        "/PIXIS/",
        "/PIXIS/Probe_Daemon/Logs",
        "/PIXIS/Probe_eth0/Logs",
        "/PIXIS/Probe_eth1/Logs",
        "/PIXIS/Probe_eth2/Logs",
        "/PIXIS/Probe_eth3/Logs"
    ]
    normalized = os.path.normpath(path).replace('\\', '/')
    if normalized in allowed_bases:
        return jsonify({"success": False, "message": "此目錄為系統基礎 Logs 或根目錄，禁止刪除！"}), 400
        
    servers = read_servers()
    server = next((s for s in servers if s["id"] == server_id), None)
    if not server:
        return jsonify({"success": False, "message": "找不到該伺服器"}), 404
        
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    
    port = int(server.get("port", 22))
    pwd = decrypt_password(server.get("password", ""))
    key_path = server.get("key_path")
    
    sftp = None
    try:
        if key_path and os.path.exists(key_path):
            key = None
            try: key = paramiko.RSAKey.from_private_key_file(key_path)
            except:
                try: key = paramiko.Ed25519Key.from_private_key_file(key_path)
                except: key = paramiko.pkey.PKey.from_private_key_file(key_path)
            ssh.connect(server.get("host"), port=port, username=server.get("username"), pkey=key, timeout=10, disabled_algorithms=None)
        else:
            ssh.connect(server.get("host"), port=port, username=server.get("username"), password=pwd, timeout=10, disabled_algorithms=None)
            
        sftp = ssh.open_sftp()
        
        try:
            stat_info = sftp.stat(path)
            import stat
            if stat.S_ISDIR(stat_info.st_mode):
                def sftp_rmtree(sftp_conn, remote_dir):
                    for attr in sftp_conn.listdir_attr(remote_dir):
                        entry = f"{remote_dir}/{attr.filename}"
                        if stat.S_ISDIR(attr.st_mode):
                            sftp_rmtree(sftp_conn, entry)
                        else:
                            sftp_conn.remove(entry)
                    sftp_conn.rmdir(remote_dir)
                sftp_rmtree(sftp, path)
            else:
                sftp.remove(path)
        except FileNotFoundError:
            return jsonify({"success": True, "message": "項目本不存在，無需刪除"}), 200
            
        return jsonify({"success": True, "message": "刪除成功！"})
    except Exception as e:
        return jsonify({"success": False, "message": f"刪除失敗: {str(e)}"}), 500
    finally:
        if sftp:
            try: sftp.close()
            except: pass
        try: ssh.close()
        except: pass

# --- 心跳自毀監控 ---
import time
import threading
import os

last_heartbeat_time = time.time()
heartbeat_lock = threading.Lock()
start_time = time.time()

@app.route('/api/heartbeat', methods=['POST'])
def heartbeat():
    global last_heartbeat_time
    with heartbeat_lock:
        last_heartbeat_time = time.time()
    return jsonify({"status": "ok"})

def monitor_heartbeat():
    # 初始寬限時間，等待瀏覽器啟動並載入頁面
    time.sleep(15)
    while True:
        time.sleep(2)
        with heartbeat_lock:
            current_time = time.time()
            elapsed_since_start = current_time - start_time
            # 若已經過寬限期且超過 300 秒未收到心跳，則自毀 (避免分頁在後台被掛起時誤殺)
            if elapsed_since_start > 15 and (current_time - last_heartbeat_time) > 300:
                print("超過 300 秒未收到心跳，背景程式自動結束。")
                os._exit(0)

# 背景開啟瀏覽器
def open_browser():

    time.sleep(1.5)
    webbrowser.open("http://localhost:7999")

if __name__ == '__main__':
    init_db()
    init_probe_test_devices()  # 確保 Probe 測試設備存在於 DB
    # 僅在非重載模式下啟動瀏覽器與監控線程
    if os.environ.get("WERKZEUG_RUN_MAIN") != "true":
        threading.Thread(target=open_browser, daemon=True).start()
        threading.Thread(target=monitor_heartbeat, daemon=True).start()
    app.run(host='0.0.0.0', port=7999, debug=False)
